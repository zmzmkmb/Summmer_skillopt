from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r'E:\桌面\暑期实训\SummerSkillOpt\reports\ACL2027_实验阶段性总结与研究方向调整报告.docx')
BLUE='2E74B5'; DARK='1F4D78'; INK='0B2545'; MUTED='5B6573'; LIGHT='F2F4F7'; CALLOUT='F4F6F9'; GOLD='7A5A00'; RED='9B1C1C'

def runfmt(r,size=11,color='000000',bold=False,italic=False):
    r.font.name='Calibri'; r._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); r.bold=bold; r.italic=italic
def pstyle(p,before=0,after=6,line=1.1,align=None):
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=line
    if align is not None: p.alignment=align
def shade(c,fill):
    tcPr=c._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def margins(c):
    tcPr=c._tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for k,v in [('top',80),('start',120),('bottom',80),('end',120)]:
        x=mar.find(qn('w:'+k))
        if x is None: x=OxmlElement('w:'+k); mar.append(x)
        x.set(qn('w:w'),str(v)); x.set(qn('w:type'),'dxa')
def geom(t,widths):
    t.autofit=False; pr=t._tbl.tblPr
    for tag,attrs in [('tblW',{'w':str(sum(widths)),'type':'dxa'}),('tblInd',{'w':'120','type':'dxa'})]:
        x=pr.find(qn('w:'+tag))
        if x is None: x=OxmlElement('w:'+tag); pr.append(x)
        for k,v in attrs.items(): x.set(qn('w:'+k),v)
    grid=t._tbl.tblGrid
    for x in list(grid): grid.remove(x)
    for w in widths: x=OxmlElement('w:gridCol'); x.set(qn('w:w'),str(w)); grid.append(x)
    for row in t.rows:
        for i,c in enumerate(row.cells):
            tcPr=c._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
            if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'),str(widths[i])); tcW.set(qn('w:type'),'dxa'); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def text(doc,s,after=6):
    p=doc.add_paragraph(); pstyle(p,0,after,1.1); runfmt(p.add_run(s)); return p
def heading(doc,s,n=1):
    p=doc.add_paragraph(style=f'Heading {n}'); pstyle(p,{1:16,2:12,3:8}[n],{1:8,2:6,3:4}[n],1.1); runfmt(p.add_run(s),{1:16,2:13,3:12}[n],BLUE if n<3 else DARK,True); return p
def bullet(doc,s):
    p=doc.add_paragraph(style='List Bullet'); pstyle(p,0,5,1.167); runfmt(p.add_run(s)); return p
def number(doc,s):
    p=doc.add_paragraph(style='List Number'); pstyle(p,0,5,1.167); runfmt(p.add_run(s)); return p
def callout(doc,label,s,fill=CALLOUT,color=INK):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; geom(t,[9360]); c=t.cell(0,0); shade(c,fill); p=c.paragraphs[0]; pstyle(p,2,2,1.1); runfmt(p.add_run(label+'：'),11,color,True); runfmt(p.add_run(s),11,color); doc.add_paragraph().paragraph_format.space_after=Pt(1)
def table(doc,heads,rows,widths,fs=10):
    t=doc.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; geom(t,widths)
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i]; shade(c,LIGHT); p=c.paragraphs[0]; pstyle(p,0,0,1.0,WD_ALIGN_PARAGRAPH.CENTER); runfmt(p.add_run(h),fs,INK,True)
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):
            p=cs[i].paragraphs[0]; pstyle(p,0,0,1.0,WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT); runfmt(p.add_run(str(v)),fs)
    doc.add_paragraph().paragraph_format.space_after=Pt(1); return t
def setup(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.492); sec.footer_distance=Inches(.492)
    for name,size,col,bef,aft in [('Normal',11,'000000',0,6),('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',12,DARK,8,4)]:
        s=doc.styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(col); s.font.bold=name!='Normal'; s.paragraph_format.space_before=Pt(bef); s.paragraph_format.space_after=Pt(aft); s.paragraph_format.line_spacing=1.1
    for name in ['List Bullet','List Number']:
        s=doc.styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(11); s.paragraph_format.space_after=Pt(5); s.paragraph_format.line_spacing=1.167
    h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pstyle(h,0,0,1.0); runfmt(h.add_run('ACL 2027 实验阶段性汇报'),9,MUTED)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; pstyle(f,0,0,1.0); runfmt(f.add_run('SummerSkillOpt · 2026年8月20日'),9,MUTED)

def build():
    OUT.parent.mkdir(parents=True,exist_ok=True); d=Document(); setup(d)
    p=d.add_paragraph(); pstyle(p,8,3,1.0); runfmt(p.add_run('ACL 2027 实验阶段性总结与研究方向调整报告'),23,INK,True)
    p=d.add_paragraph(); pstyle(p,0,12,1.0); runfmt(p.add_run('——从“typed prior 是否直接提升效果”转向“activation–grounding–utility gap”'),13,MUTED)
    for a,b in [('汇报对象','导师阶段性沟通'),('项目','SummerSkillOpt / ACL 2027'),('报告日期','2026年8月20日'),('当前状态','Phase 4C-R1 已完成闭环；后续 provider 调用关闭')]:
        p=d.add_paragraph(); pstyle(p,0,2,1.0); runfmt(p.add_run(a+'：'),10.5,INK,True); runfmt(p.add_run(b),10.5)
    callout(d,'先说结论','目前实验结果不支持把“历史 skill prior 能稳定提升真实任务最终答案”作为论文主结论。比较稳定的证据是：prior 可以被模型激活，合同和证据 ID 可以被修复，但“被采用”不自动等于可验证的中间操作，更不稳定地等于最终答案收益。因此，下一阶段应转向 activation、grounding 和 utility 的分解，并提出 prior admission / abstention 方法。','E8EEF5')

    heading(d,'1. 报告目的与当前判断'); text(d,'这份报告不是把已有结果包装成“已经证明了原假设”，而是给出一个可以和老师直接讨论的阶段性判断：协议、合同、证据 ID、held-out 运行和审计链条已经做得比较完整，但原始主效应没有被稳定识别出来。继续重复同类请求，最多增加样本，不能自动修复研究问题本身。')
    text(d,'项目仓库当前记录为 51 个阶段、6206 个不可变运行，当前阶段为 Phase 4C-R1。按照项目协议，已完成运行和账本保持不可变，provider/API 调用处于关闭状态。本报告只做研究总结和下一阶段设计，不启动新的外部实验。')
    heading(d,'2. 原始研究假设：一开始想证明什么'); text(d,'原始主张可以概括为：如果历史成功经验被整理成 typed / scoped skill priors，并在持续学习或路由过程中提供给模型，那么这些 prior 应稳定提高真实任务的可靠性、答案正确率或单位 token 效率。这个主张隐含了多个前提：模型会看到 prior；会采用正确 prior；采用后会改变正确的中间操作；中间操作会传递到最终答案；收益不是来自额外上下文、提示格式或任务难度差异。前几轮实验说明，这些前提不能继续被当成一个黑箱。')

    heading(d,'3. 已完成实验与主要结果'); heading(d,'3.1 Phase 2：真实任务效果没有形成稳定证据',2); text(d,'Phase 2 的 typed-prior gate 总体为 inconclusive，v25 replication 给出负结果。我们没有得到足够稳定、可以支撑“typed prior 在真实任务上普遍有效”的结果。这不等于 prior 一定无效，而是说明原始主效应在当前任务、协议和样本下没有被可靠识别。')
    heading(d,'3.2 Phase 3：模型会采用 prior，但 activation 不等于 utility',2); text(d,'Phase 3B 支持机制层面的采用：contextual typed bundle 能被模型识别和采用。但 Phase 3C、3F 显示，采用并不稳定转化为中间操作，更不稳定地改变最终答案。最值得保留的洞见是：activation 与 utility 之间存在断层。')
    table(d,['阶段','主要观察','可以支持的表述'],[['Phase 2','真实任务 gate inconclusive；replication 为负','不能宣称稳定的 typed-prior 主效应'],['Phase 3B','contextual typed bundle 采用率明显','prior activation 机制存在'],['Phase 3C/3F','操作变化多于答案变化；中介 gate inconclusive','activation 不等于 grounding / utility'],['Phase 3H','答案 grounding gate inconclusive','需要直接测量 prior 是否改变可验证答案依据']],[1500,3900,3960])
    heading(d,'3.3 Phase 4B：修复的是可见性与协议，不是因果收益',2); text(d,'Phase 4B 发现，原先 provider 没有看到完整六字段合同，evidence ID 也没有真正进入模型可见消息。修复后，合同和 evidence ID 可见性改善，100 行 calibration 全部 contract-valid 且 evidence-ID resolving，contextual selection 为 1.0。这个结果证明的是“协议和传输可以修好”，不是“typed prior 已经带来因果收益”。另外，global_only 与 contextual_typed 的 20 对 payload 仍 transport-identical，不能用来做有效条件比较。')
    heading(d,'3.4 Phase 4C：技术执行基本通过，但覆盖不足以给出正式 gate',2); text(d,'Phase 4C-R1 修复了 held-out 设计：80 个任务、400 行、5 个条件，每行嵌入六字段合同并标注 evidence ID，400 个 transport payload 唯一，且与已花费 Phase 4B 身份不重叠。实际启动 396 次，完成 395 条响应，1 次 terminal hard-stop；缺失 5 行全部属于同一个 task。已知阶段成本 CNY 2.992020，总 token 1,196,202，retries 为 0，pacing 与 hash-chain 审计通过，授权自动关闭。')
    table(d,['条件','正确数 / 计分数','准确率','合同有效 / evidence ID'],[['cold','73 / 79','92.41%','79 / 79；79 / 79'],['global_only','78 / 79','98.73%','79 / 79；79 / 79'],['contextual_typed','78 / 79','98.73%','79 / 79；79 / 79'],['shuffled_typed','78 / 79','98.73%','79 / 79；79 / 79'],['incompatible_control','76 / 79','96.20%','79 / 79；79 / 79']],[2200,2100,1500,3560],9.5)
    text(d,'这些数字只能做描述性诊断，不能写成 400 行 frozen gate 的正式正面或负面结论：400 行覆盖没有完成；累计成本账本此前存在需要修正的历史 Phase 4B 成本带入问题。因此旧 ledger、closure 和 terminal attempt 必须保持不可变，不能简单“把上限调高后继续跑”。')
    heading(d,'3.5 Phase 4C paired diagnostic：差异很小，且没有回答因果问题',2); text(d,'在 79 个完整 task grid 上，contextual_typed 与 global_only 是 0 wins、0 losses、79 ties，0 个 answer changes；与 cold 比是 5 wins、0 losses、74 ties，6 个 answer changes；与 incompatible_control 比是 2 wins、0 losses、77 ties，3 个 answer changes。即使出现少量差异，也不足以说明 typed prior 本身造成了差异。')

    heading(d,'4. 当前最重要的阶段性发现'); callout(d,'核心发现','目前最可信的不是“prior 有效”或“prior 无效”，而是：prior 的 activation、evidence grounding 和最终 utility 是三个不同环节；前一个环节通过，不代表后两个环节会自动通过。','FFF8E8',GOLD); text(d,'这也解释了为什么会出现“模型明明用了 prior，但答案没有明显变好”：模型可以在语言层面复述、引用或选择一个 prior，同时仍然没有把它接到正确的实体、关系、计算或证据链上。若论文只报告“是否采用 prior”，就会把机制信号误读成任务收益。')
    heading(d,'5. 为什么需要调整论文方向'); text(d,'继续沿原方向堆更多 API 调用，最大的风险不是成本，而是科学问题没有变清楚。即使再跑一批，看到平均准确率轻微波动，仍很难回答收益来自 prior 内容、提示长度、合同修复、模型随机性，还是任务本身可解性。ACL 审稿人很可能会把它看成“合理但证据边界模糊的框架”。')
    text(d,'因此方向调整不是把负结果改写成正结果，而是把已有失败和不确定性转化成新问题：什么时候允许一个 prior 进入上下文？什么时候要求它先通过 evidence-grounded admission？什么时候模型应该 abstain 或 fallback，而不是强行使用看起来相关的历史经验？')
    table(d,['原方向','调整后的方向'],[['证明 typed/scoped prior 平均提升最终答案','解释 activation–grounding–utility 三段链路为何脱钩'],['默认 prior 应被使用','prior 经过 admission；无法改变可验证中间结果时降权或拒绝'],['主要指标是最终 accuracy','同时测 activation、grounding、answer change、utility 与 abstention'],['继续增加同类 provider 请求','先做零网络设计、反事实任务和停止条件，再决定是否执行']],[4680,4680])

    heading(d,'6. 新研究主线：activation–grounding–utility gap'); text(d,'新的工作假设是：skill priors 可能可靠地被模型激活，但 activation 只有在被正确 grounding 到任务证据和可验证中间操作之后，才可能转化为 utility；因此，可靠 continual system 需要显式的 prior admission、grounding 检查和 abstention / fallback 机制。')
    heading(d,'6.1 三个可分离的测量层',2); bullet(d,'Activation：模型是否选中了 prior、是否识别出适用范围、是否拒绝明显不相关 prior。'); bullet(d,'Grounding：prior 是否改变可验证的实体、属性、桥接关系、计算步骤或证据引用，而不是只改变表面措辞。'); bullet(d,'Utility：grounding 后是否提高最终答案正确率、减少错误、降低 token 成本，或在不确定时更早 abstain。')
    heading(d,'6.2 方法贡献的候选形态',2); bullet(d,'Prior admission：把 prior 放进主上下文前，先检查任务类型、适用范围、证据 ID 和反事实可用性。'); bullet(d,'Evidence-grounded selection：要求 prior 绑定可解析证据片段，并输出可审计中间操作。'); bullet(d,'Abstention / fallback：当 prior 无法改变或支持可验证中间结果时，自动降权、拒绝或回退到 cold / global-only。'); bullet(d,'Counterfactual evaluation：构造“使用 prior 应该改变答案”的任务，避免只测到表面采用。')

    heading(d,'7. 下一阶段拟解决的问题'); number(d,'把 prior admission 定义成明确决策问题：输入是任务、候选 prior 和证据；输出是 admit、abstain 或 fallback。'); number(d,'为每个任务预注册 answer-sensitive 中间变量，例如实体桥接、属性比较、数值计算或证据支持关系。'); number(d,'设置 no prior、always-use prior、random / shuffled prior、wrong / incompatible prior 等基线，再与新 admission 方法比较。'); number(d,'把“操作变化但答案不变”单独作为 failure mode 统计，而不是混入总体 accuracy。'); number(d,'先做零网络 preflight，检查条件是否 transport-distinct、证据是否真正可见、任务是否能产生反事实答案变化。')
    heading(d,'8. 后续实验计划与停止条件'); table(d,['阶段','目标','关键产物','停止条件'],[['A. 设计','冻结 admission / abstention 任务与反事实标签','零网络 schedule、fingerprint、分析计划','条件不可识别或答案不敏感则重设计'],['B. 本地审计','验证合同、evidence、payload distinctness','preflight report、hash-chain 检查','任一审计失败，不进入 live'],['C. 小规模校准','验证协议和基本可解性','有限授权校准分析','contract / capability gate 不通过则终止'],['D. held-out','比较 always-use 与 admission / abstention','成对 task-grid 分析','首个失败、预算或覆盖条件触发即闭环']],[1500,3000,2700,2160],9.5)
    callout(d,'项目纪律','下一阶段不以“多跑一些请求”为默认动作。任何新 provider 调用都必须有独立 zero-network preflight、新 fingerprint、明确预算和新的精确授权；已完成 ledger、closure、terminal attempt 和成本记录不回写、不覆盖。','FDECEC',RED)
    heading(d,'9. 风险与预期贡献边界'); text(d,'新方法可能仍只改善 abstention 或中间操作，而不显著提高平均 accuracy；有些任务也许本身对 prior 不敏感；模型还可能学会“看起来 grounding”但无法完成真正关系推理。因此论文不应承诺普遍 accuracy 增益，而应聚焦于更可靠的诊断与控制：识别 activation–utility gap，给出可复现 admission / abstention 机制，并说明它什么时候有效、什么时候应该拒绝 prior。')
    text(d,'如果后续仍没有正面的最终答案收益，只要失配模式被稳定复现，并且 admission 能减少错误使用 prior 或降低无效上下文成本，仍可形成一篇逻辑完整的分析型方法论文。')
    heading(d,'10. 给导师的口头总结（可直接照着说）'); callout(d,'口头版','老师，我们前面几轮实验把协议和运行链条做扎实了，但结果没有支持最初那个强主张：不能说 typed historical prior 能稳定提升真实任务最终答案。Phase 3 反而给了更清楚的信号——模型会采用 prior，但采用不等于它接到了正确证据和中间推理上。Phase 4B/4C 又说明，合同和 evidence 可见性可以修复，但修复协议也不等于因果收益。所以下一步不再重点堆 API 请求，而是研究 activation–grounding–utility gap，做 prior admission、evidence-grounded selection 和 abstention/fallback。这样论文贡献会从“一个框架平均有效”转向“解释什么时候 prior 真正有用，以及什么时候应该拒绝它”。')
    heading(d,'附录：本报告引用的关键状态'); text(d,'实验状态来源：paper/acl2027/experiment_state.json 与 paper/acl2027/CROSS_CONVERSATION_PROTOCOL.md。Phase 4C 描述性分析来源：artifacts/acl2027_phase4c_repaired_heldout_live_v2/phase4c_analysis_v2.json。报告生成日期：2026年8月20日。')
    d.save(OUT); print(OUT)

if __name__=='__main__': build()
