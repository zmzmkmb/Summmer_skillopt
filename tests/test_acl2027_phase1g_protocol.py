from __future__ import annotations

import json

import pytest
from openpyxl import Workbook, load_workbook

from scripts.validate_acl2027_phase1g_protocol import CONFIG_PATH, audit
from skillopt.evaluation.cross_task_protocol import (
    ProtocolError,
    apply_spreadsheet_edits,
    parse_json_object_preserving_raw,
    validate_officeqa_trace,
    validate_spreadsheet_edits,
)


def test_whitespace_key_normalization_preserves_audit():
    result = parse_json_object_preserving_raw('{"answer":"x"," evidence":"table"}')
    assert result.value["evidence"] == "table"
    assert result.normalizations == ({"raw_key": " evidence", "normalized_key": "evidence"},)
    assert result.raw_text.startswith("{")


def test_normalization_rejects_collisions():
    with pytest.raises(ProtocolError, match="collision"):
        parse_json_object_preserving_raw('{"answer":1," answer":2}')


def test_calendar_trace_blocks_fiscal_shortcut():
    payload = {
        "answer": "1580",
        "evidence": [{"source_path": "x", "locator": "row", "period": "FY1940", "value": 1580}],
        "calculation": {"operation": "sum", "time_basis": "fiscal", "operands": [{"period": "FY1940", "value": 1580}], "result": 1580},
    }
    result = validate_officeqa_trace(payload, required_time_basis="calendar", expected_period_count=12)
    assert result["retrieval_supported"] is True
    assert result["aggregation_valid"] is False
    assert result["eligible_for_gate"] is False


def test_spreadsheet_edits_require_exact_coverage_and_persist(tmp_path):
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.save(input_path)
    workbook.close()
    payload = {
        "target_range": "B2:C2",
        "edits": [
            {"sheet": "Sheet1", "cell": "B2", "formula": "=1+1"},
            {"sheet": "Sheet1", "cell": "C2", "formula": "=2+2"},
        ],
        "explanation": "two formulas",
    }
    assert validate_spreadsheet_edits(payload, expected_target_range="B2:C2")["exact_coverage"] is True
    apply_spreadsheet_edits(input_path, output_path, payload)
    reopened = load_workbook(output_path, data_only=False)
    assert reopened["Sheet1"]["B2"].value == "=1+1"
    reopened.close()


def test_phase1g_full_local_audit_passes(tmp_path):
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8-sig"))
    result = audit(config, tmp_path)
    assert result["decision"] == "local_protocol_passed_paid_smoke_still_closed"
    assert result["officeqa_strata"]["difficulty"] == {"easy": 4, "hard": 4}
    assert result["network_calls"] == result["paid_api_calls"] == 0