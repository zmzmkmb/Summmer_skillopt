"""Auditable response contracts for ACL 2027 cross-task validation."""
from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


class ProtocolError(ValueError):
    """Raised when a model response cannot satisfy the frozen protocol."""


@dataclass(frozen=True)
class NormalizedObject:
    raw_text: str
    value: dict[str, Any]
    normalizations: tuple[dict[str, str], ...]


def parse_json_object_preserving_raw(text: str) -> NormalizedObject:
    """Parse one JSON object and trim key whitespace without hiding changes."""
    cleaned = text.strip()
    fence = chr(96) * 3
    if cleaned.startswith(fence):
        cleaned = re.sub(r"^" + re.escape(fence) + r"(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\s*" + re.escape(fence) + r"$", "", cleaned)
    parsed = json.loads(cleaned)
    if not isinstance(parsed, dict):
        raise ProtocolError("response must be a JSON object")
    normalized: dict[str, Any] = {}
    changes: list[dict[str, str]] = []
    for raw_key, value in parsed.items():
        key = str(raw_key).strip()
        if not key:
            raise ProtocolError("response contains an empty key after normalization")
        if key in normalized:
            raise ProtocolError(f"key collision after normalization: {key}")
        normalized[key] = value
        if key != raw_key:
            changes.append({"raw_key": str(raw_key), "normalized_key": key})
    return NormalizedObject(text, normalized, tuple(changes))


def validate_officeqa_trace(payload: dict[str, Any], *, required_time_basis: str, expected_period_count: int | None) -> dict[str, Any]:
    """Validate evidence provenance separately from temporal aggregation."""
    missing = [key for key in ("answer", "evidence", "calculation") if key not in payload]
    if missing:
        raise ProtocolError(f"OfficeQA response omitted required fields: {missing}")
    evidence = payload["evidence"]
    calculation = payload["calculation"]
    if not isinstance(evidence, list) or not evidence:
        raise ProtocolError("OfficeQA evidence must be a non-empty list")
    if not isinstance(calculation, dict):
        raise ProtocolError("OfficeQA calculation must be an object")
    evidence_fields = ("source_path", "locator", "period", "value")
    retrieval_supported = all(isinstance(item, dict) and all(str(item.get(key, "")).strip() for key in evidence_fields) for item in evidence)
    time_basis_match = calculation.get("time_basis") == required_time_basis
    operands = calculation.get("operands")
    operand_periods = {str(item.get("period")) for item in operands or [] if isinstance(item, dict) and item.get("period") is not None}
    period_count_match = expected_period_count is None or len(operand_periods) == expected_period_count
    aggregation_valid = bool(str(calculation.get("operation", "")).strip() and time_basis_match and period_count_match and calculation.get("result") is not None)
    return {"retrieval_supported": retrieval_supported, "time_basis_match": time_basis_match, "period_count_match": period_count_match, "observed_period_count": len(operand_periods), "aggregation_valid": aggregation_valid, "eligible_for_gate": retrieval_supported and aggregation_valid}


def target_cells(target_range: str) -> list[str]:
    """Expand an A1 range in row-major order."""
    bare = target_range.split("!", 1)[-1].strip().strip("'\"")
    min_col, min_row, max_col, max_row = range_boundaries(bare)
    return [f"{get_column_letter(col)}{row}" for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1)]


def validate_spreadsheet_edits(payload: dict[str, Any], *, expected_target_range: str) -> dict[str, Any]:
    missing = [key for key in ("target_range", "edits", "explanation") if key not in payload]
    if missing:
        raise ProtocolError(f"SpreadsheetBench response omitted required fields: {missing}")
    if str(payload["target_range"]).strip() != expected_target_range:
        raise ProtocolError("SpreadsheetBench target range does not match the task")
    edits = payload["edits"]
    if not isinstance(edits, list) or not edits:
        raise ProtocolError("SpreadsheetBench edits must be a non-empty list")
    cells = [str(item.get("cell", "")).upper() for item in edits if isinstance(item, dict)]
    expected = target_cells(expected_target_range)
    if len(cells) != len(set(cells)):
        raise ProtocolError("SpreadsheetBench edits contain duplicate cells")
    if set(cells) != set(expected):
        raise ProtocolError("SpreadsheetBench edits do not exactly cover the target range")
    for item in edits:
        if not all(str(item.get(key, "")).strip() for key in ("sheet", "cell", "formula")):
            raise ProtocolError("SpreadsheetBench edit omitted sheet, cell, or formula")
        if not str(item["formula"]).startswith("="):
            raise ProtocolError("SpreadsheetBench formulas must start with '='")
    return {"target_cells": len(expected), "exact_coverage": True, "formulas_valid": True}


def apply_spreadsheet_edits(input_path: Path, output_path: Path, payload: dict[str, Any]) -> None:
    """Materialize per-cell edits into a fresh workbook and verify persistence."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)
    workbook = load_workbook(output_path, data_only=False)
    for edit in payload["edits"]:
        sheet = str(edit["sheet"])
        if sheet not in workbook.sheetnames:
            workbook.close()
            raise ProtocolError(f"worksheet not found: {sheet}")
        workbook[sheet][str(edit["cell"])] = str(edit["formula"])
    workbook.save(output_path)
    workbook.close()
    reopened = load_workbook(output_path, data_only=False, read_only=True)
    try:
        for edit in payload["edits"]:
            actual = reopened[str(edit["sheet"])][str(edit["cell"])].value
            if actual != str(edit["formula"]):
                raise ProtocolError(f"formula did not persist at {edit['sheet']}!{edit['cell']}")
    finally:
        reopened.close()