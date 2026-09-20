#!/usr/bin/env python3
"""Replay immutable Phase 1I responses through stricter zero-network diagnostics."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "configs/acl2027/phase1j_no_paid_capability_calibration_v1.json"
DEFAULT_OUTPUT = ROOT / "artifacts/acl2027_phase1j_no_paid_capability_calibration_v1"
EDIT_OBJECT_RE = re.compile(
    r'\{\s*"sheet"\s*:\s*"(?:\\.|[^"\\])*"\s*,\s*'
    r'"cell"\s*:\s*"(?:\\.|[^"\\])*"\s*,\s*'
    r'"formula"\s*:\s*"(?:\\.|[^"\\])*"\s*\}',
    re.DOTALL,
)


class CalibrationError(RuntimeError):
    """Raised when frozen calibration inputs or outputs are invalid."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, indent=2, ensure_ascii=True, sort_keys=True) + "\n", encoding="utf-8")


def manifest_path(path: Path) -> str:
    try:
        value = path.relative_to(ROOT)
    except ValueError:
        value = path
    return str(value).replace("\\", "/")


def as_decimal(value: Any) -> Decimal | None:
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def numeric_mentions(text: Any) -> set[Decimal]:
    values: set[Decimal] = set()
    for match in re.findall(r"(?<![A-Za-z])[-+]?\d[\d,]*(?:\.\d+)?", str(text)):
        value = as_decimal(match)
        if value is not None:
            values.add(value)
    return values


def load_inputs(config_path: Path = DEFAULT_CONFIG) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    execution = config["execution"]
    if execution["network_calls_allowed"] or execution["provider_calls"] != 0 or execution["paid_api_allowed"]:
        raise CalibrationError("Phase 1J must remain zero-network and no-paid")
    inputs = config["immutable_inputs"]
    for path_key, hash_key in (
        ("phase1i_results", "phase1i_results_sha256"),
        ("phase1i_scientific_results", "phase1i_scientific_results_sha256"),
    ):
        path = ROOT / inputs[path_key]
        if sha256_file(path) != inputs[hash_key]:
            raise CalibrationError(f"immutable input hash mismatch: {inputs[path_key]}")
    rows = [
        json.loads(line)
        for line in (ROOT / inputs["phase1i_results"]).read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    identities = [(row.get("task_family"), row.get("task_id")) for row in rows]
    if identities != [("OfficeQA", "UID0001"), ("SpreadsheetBench", "45635")]:
        raise CalibrationError("Phase 1I response identity or order changed")
    return config, rows


def diagnose_officeqa(config: dict[str, Any], record: dict[str, Any]) -> dict[str, Any]:
    payload = record.get("parsed_response") or {}
    evidence = payload.get("evidence") if isinstance(payload, dict) else None
    calculation = payload.get("calculation") if isinstance(payload, dict) else None
    evidence = evidence if isinstance(evidence, list) else []
    calculation = calculation if isinstance(calculation, dict) else {}
    expected_count = int(config["officeqa"]["expected_period_count"])
    reference = as_decimal(config["officeqa"]["reference_answer"])

    evidence_items_valid = all(
        isinstance(item, dict)
        and all(str(item.get(key, "")).strip() for key in ("source_path", "locator", "period", "value"))
        and as_decimal(item.get("value")) is not None
        for item in evidence
    )
    periods = [str(item.get("period")) for item in evidence if isinstance(item, dict)]
    values = [as_decimal(item.get("value")) for item in evidence if isinstance(item, dict)]
    evidence_sum = sum((value for value in values if value is not None), Decimal(0))
    retrieval_failure = not (
        evidence_items_valid
        and len(evidence) == expected_count
        and len(set(periods)) == expected_count
        and evidence_sum == reference
    )

    operands = calculation.get("operands")
    operand_items = operands if isinstance(operands, list) else []
    operand_schema_valid = (
        len(operand_items) == expected_count
        and all(
            isinstance(item, dict)
            and set(item) == {"period", "value"}
            and str(item.get("period", "")).strip()
            and as_decimal(item.get("value")) is not None
            for item in operand_items
        )
    )
    if operand_schema_valid:
        operand_pairs = {(str(item["period"]), as_decimal(item["value"])) for item in operand_items}
        evidence_pairs = {(str(item["period"]), as_decimal(item["value"])) for item in evidence}
        operand_schema_valid = operand_pairs == evidence_pairs

    reported_result = as_decimal(calculation.get("result"))
    channels = {
        "retrieval_or_provenance_failure": retrieval_failure,
        "operand_schema_failure": not operand_schema_valid,
        "arithmetic_failure": reported_result != evidence_sum,
        "answer_mismatch": evidence_sum not in numeric_mentions(payload.get("answer", "")),
    }
    return {
        "task_family": "OfficeQA",
        "task_id": record["task_id"],
        "raw_phase1i_task_valid": record.get("task_valid") is True,
        "evidence_count": len(evidence),
        "unique_evidence_periods": len(set(periods)),
        "deterministic_evidence_sum": int(evidence_sum),
        "reference_answer": int(reference) if reference is not None else None,
        "reported_calculation_result": int(reported_result) if reported_result is not None else None,
        "diagnostic_channels": channels,
        "diagnostic_recovery_task_success": False,
    }


def recover_complete_edits(raw_text: str) -> list[dict[str, Any]]:
    edits = []
    for match in EDIT_OBJECT_RE.finditer(raw_text):
        value = json.loads(match.group(0))
        if set(value) == {"sheet", "cell", "formula"}:
            edits.append(value)
    return edits


def golden_edits(config: dict[str, Any]) -> list[dict[str, str]]:
    path = ROOT / config["spreadsheetbench"]["golden_workbook"]
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    edits = []
    for row in range(5, 8):
        for column in range(2, 6):
            cell = sheet.cell(row=row, column=column)
            edits.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": str(cell.value)})
    workbook.close()
    return edits


def diagnose_spreadsheet(config: dict[str, Any], record: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    raw_text = str(record.get("raw_text", ""))
    try:
        parsed = json.loads(raw_text)
        parse_failure = False
    except json.JSONDecodeError:
        parsed = None
        parse_failure = True
    recovered = recover_complete_edits(raw_text)
    expected = golden_edits(config)
    expected_by_cell = {item["cell"]: item["formula"] for item in expected}
    recovered_cells = [str(item["cell"]).upper() for item in recovered]
    exact_coverage = (
        len(recovered_cells) == 12
        and len(set(recovered_cells)) == 12
        and set(recovered_cells) == set(expected_by_cell)
    )
    mismatches = [
        {"cell": item["cell"], "predicted": item["formula"], "expected": expected_by_cell.get(str(item["cell"]).upper())}
        for item in recovered
        if expected_by_cell.get(str(item["cell"]).upper()) != item["formula"]
    ]
    schema_failure = not (
        isinstance(parsed, dict)
        and set(parsed) == {"target_range", "edits", "explanation"}
        and parsed.get("target_range") == config["spreadsheetbench"]["target_range"]
        and isinstance(parsed.get("explanation"), str)
        and len(parsed["explanation"]) <= config["spreadsheetbench"]["response_contract"]["explanation_max_characters"]
        and exact_coverage
    )
    compact_payload = {
        "target_range": config["spreadsheetbench"]["target_range"],
        "edits": expected,
        "explanation": "Compare the three people independently within each source column and assign 6/3/2 points to tied maxima.",
    }
    compact_json = json.dumps(compact_payload, ensure_ascii=True, separators=(",", ":"))
    diagnosis = {
        "task_family": "SpreadsheetBench",
        "task_id": record["task_id"],
        "raw_phase1i_task_valid": record.get("task_valid") is True,
        "complete_edits_recovered_for_diagnosis": len(recovered),
        "recovered_exact_cell_coverage": exact_coverage,
        "golden_formula_mismatch_count": len(mismatches),
        "mismatches": mismatches,
        "diagnostic_channels": {
            "json_completion_failure": parse_failure,
            "edit_schema_failure": schema_failure,
            "formula_semantics_failure": len(mismatches) != 0,
        },
        "diagnostic_recovery_task_success": False,
    }
    compact = {
        "payload": compact_payload,
        "compact_json_characters": len(compact_json),
        "compact_json_utf8_bytes": len(compact_json.encode("utf-8")),
        "closes_as_json": json.loads(compact_json) == compact_payload,
        "within_6000_byte_guard": len(compact_json.encode("utf-8")) <= 6000,
    }
    return diagnosis, compact


def run_calibration(config_path: Path = DEFAULT_CONFIG, output_dir: Path = DEFAULT_OUTPUT) -> dict[str, Any]:
    if output_dir.exists():
        raise CalibrationError(f"refusing to overwrite immutable artifact: {output_dir}")
    config, rows = load_inputs(config_path)
    output_dir.mkdir(parents=True)
    office = diagnose_officeqa(config, rows[0])
    spreadsheet, compact = diagnose_spreadsheet(config, rows[1])
    contracts = {
        "officeqa": {
            "operand_item_schema": {"period": "string", "value": "number"},
            "validator_rule": "Recompute from evidence; answer and calculation.result must both equal the recomputed value. Never silently replace a mismatch.",
        },
        "spreadsheetbench": {
            "top_level_keys": ["target_range", "edits", "explanation"],
            "edit_item_schema": {"sheet": "string", "cell": "A1 reference", "formula": "Excel formula"},
            "semantics": config["spreadsheetbench"]["development_semantics"],
            "explanation_max_characters": 160,
        },
    }
    results = {
        "experiment": config["experiment"],
        "network_calls": 0,
        "paid_api_calls": 0,
        "source_phase1i_task_valid_labels_preserved": not office["raw_phase1i_task_valid"] and not spreadsheet["raw_phase1i_task_valid"],
        "officeqa": office,
        "spreadsheetbench": spreadsheet,
        "compact_spreadsheet_contract": compact,
        "model_assessment": {
            "qwen3.6_flash_currently_qualified": False,
            "reason": "Both immutable Phase 1I base tasks remain invalid after diagnosis; recovery only localizes failures.",
            "revised_contract_viability_established": False,
            "next_model_decision": "Use a separately authorized paired confirmation that includes qwen3.6-flash and a stronger model candidate; do not launch the 24/40 batch first.",
        },
        "paper_consequence": "No main-claim direction change: this is protocol/capability calibration, not prior-transfer or triage evidence. The real-task route remains blocked pending a task-valid model-and-contract floor.",
    }
    diagnostics_path = output_dir / "diagnostics.json"
    contracts_path = output_dir / "revised_response_contracts.json"
    write_json(diagnostics_path, results)
    write_json(contracts_path, contracts)
    diagnostics_hash = sha256_file(diagnostics_path)
    contracts_hash = sha256_file(contracts_path)
    manifest = {
        "schema_version": 1,
        "experiment": config["experiment"],
        "phase": "1J",
        "status": "completed",
        "config_path": str(config_path.relative_to(ROOT)).replace("\\", "/"),
        "config_sha256": sha256_file(config_path),
        "expected_runs": 2,
        "available_runs": 2,
        "complete_grid": True,
        "network_calls": 0,
        "paid_api_calls": 0,
        "decision": "diagnostics_passed_task_valid_floor_still_unestablished",
        "scientific_results_path": manifest_path(diagnostics_path),
        "aggregate_fingerprint": diagnostics_hash,
        "runs": [
            {"run_id": "phase1j_officeqa_uid0001_replay", "result_path": manifest_path(diagnostics_path), "file_sha256": diagnostics_hash},
            {"run_id": "phase1j_spreadsheetbench_45635_contract", "result_path": manifest_path(contracts_path), "file_sha256": contracts_hash}
        ]
    }
    write_json(output_dir / "run_manifest.json", manifest)
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(run_calibration(args.config, args.output_dir), indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
