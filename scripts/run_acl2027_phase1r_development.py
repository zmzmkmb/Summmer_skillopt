#!/usr/bin/env python3
"""Materialize and safely execute the ACL 2027 Phase 1R development plan."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
from copy import deepcopy
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from skillopt.envs.officeqa.tool_runtime import (
    _iter_oracle_refs,
    _locate_parsed_json,
    _render_parsed_page,
    resolve_docs_roots,
)

DEFAULT_CONFIG = ROOT / "configs/acl2027/phase1r_development_runner_preflight_v1.json"


class RunnerError(RuntimeError):
    """Raised when a frozen runner invariant is violated."""


class HardStop(RunnerError):
    """Raised after a terminal record is persisted."""


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=True, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def stable_hash(value: Any) -> str:
    payload = json.dumps(
        value, ensure_ascii=True, sort_keys=True, separators=(",", ":")
    ).encode()
    return hashlib.sha256(payload).hexdigest()


def relative(path: Path) -> str:
    try:
        path = path.relative_to(ROOT)
    except ValueError:
        pass
    return str(path).replace("\\", "/")


def validate_config(config_path: Path = DEFAULT_CONFIG) -> dict[str, Any]:
    config = read_json(config_path)
    if config.get("phase") != "1R":
        raise RunnerError("config phase must be 1R")
    execution = config["execution"]
    closed = (
        execution["protocol_only"]
        and execution["paid_api_allowed"] is False
        and execution["network_calls_allowed"] is False
        and execution["provider_calls"] == 0
        and execution["formal_scaling_allowed"] is False
    )
    if not closed:
        raise RunnerError("Phase 1R must remain zero-network and unpaid")
    if (
        execution["max_provider_attempts_per_logical_call"] != 1
        or execution["sdk_max_retries"] != 0
        or execution["explicit_retries"] != 0
    ):
        raise RunnerError("Phase 1R retry or attempt contract changed")
    if "max_tokens" not in execution["request_body_forbidden_keys"]:
        raise RunnerError("max_tokens omission is not frozen")
    for name, item in config["source_contracts"].items():
        if not isinstance(item, dict) or "path" not in item:
            continue
        path = ROOT / item["path"]
        if not path.is_file() or sha256_file(path) != item["sha256"]:
            raise RunnerError(f"immutable source hash mismatch: {name}")
    return config


def _office_rows(path: Path) -> dict[str, dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return {
            str(row["uid"]).strip(): {key: str(value or "") for key, value in row.items()}
            for row in csv.DictReader(handle)
            if str(row.get("uid", "")).strip()
        }


def _query_terms(question: str) -> set[str]:
    stop = {
        "the", "and", "for", "from", "that", "this", "with", "were", "what",
        "which", "into", "across", "between", "calculate", "using", "return",
        "year", "years", "calendar", "million", "millions", "dollars",
    }
    return {
        token for token in re.findall(r"[a-z0-9]+", question.lower())
        if len(token) >= 3 and token not in stop
    }


def compact_office_evidence(
    row: dict[str, str], docs_roots: list[str], max_chars: int
) -> str:
    terms = _query_terms(row["question"])
    blocks: list[str] = []
    for source_file, page_number, source_doc in _iter_oracle_refs(
        row["source_files"], row["source_docs"]
    ):
        parsed = _locate_parsed_json(source_file, docs_roots)
        if parsed is None:
            continue
        text = _render_parsed_page(str(parsed), page_number)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        scores = [
            (sum(term in line.lower() for term in terms), index)
            for index, line in enumerate(lines)
        ]
        selected: set[int] = set()
        for score, index in sorted(scores, key=lambda item: (-item[0], item[1])):
            if score <= 0 and selected:
                break
            selected.update(range(max(0, index - 2), min(len(lines), index + 3)))
            if sum(len(lines[i]) + 1 for i in selected) >= max_chars // max(
                1, len(_iter_oracle_refs(row["source_files"], row["source_docs"]))
            ):
                break
        if not selected:
            selected.update(range(min(20, len(lines))))
        excerpt = "\n".join(lines[index] for index in sorted(selected))
        blocks.append(
            f"### {source_file} page {page_number}\n"
            f"Source URL: {source_doc}\n{excerpt}"
        )
    joined = "\n\n".join(blocks)
    return joined[:max_chars].rstrip()


def _json_cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    formula_text = getattr(value, "text", None)
    if formula_text is not None:
        return {
            "type": type(value).__name__,
            "formula": str(formula_text),
            "ref": str(getattr(value, "ref", "")),
        }
    return {"type": type(value).__name__, "value": str(value)}


def _sheet_cells(sheet: Any, rows: list[int] | None = None) -> dict[str, Any]:
    selected = set(rows) if rows is not None else None
    cells: dict[str, Any] = {}
    for row_index, row in enumerate(sheet.iter_rows(), start=1):
        if selected is not None and row_index not in selected:
            continue
        for cell in row:
            if cell.value is not None:
                cells[cell.coordinate] = _json_cell_value(cell.value)
    return cells


def compact_workbook_snapshot(
    workbook_path: Path,
    *,
    max_nonempty: int,
    head_rows: int,
    tail_rows: int,
) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        sheet_counts = {
            sheet.title: sum(
                1 for row in sheet.iter_rows() for cell in row if cell.value is not None
            )
            for sheet in workbook.worksheets
        }
        total = sum(sheet_counts.values())
        sheets: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            truncated = total > max_nonempty
            rows = None
            if truncated:
                rows = list(range(1, min(sheet.max_row, head_rows) + 1))
                rows += list(
                    range(max(1, sheet.max_row - tail_rows + 1), sheet.max_row + 1)
                )
                rows = sorted(set(rows))
            sheets.append(
                {
                    "sheet": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "nonempty_cells": sheet_counts[sheet.title],
                    "snapshot_truncated": truncated,
                    "cells": _sheet_cells(sheet, rows),
                }
            )
        return {"total_nonempty_cells": total, "sheets": sheets}
    finally:
        workbook.close()


def _skill_texts(config: dict[str, Any]) -> dict[str, str]:
    source = config["source_contracts"]
    return {
        "officeqa_skill": (ROOT / source["officeqa_skill"]["path"]).read_text(
            encoding="utf-8"
        ),
        "spreadsheetbench_skill": (
            ROOT / source["spreadsheetbench_skill"]["path"]
        ).read_text(encoding="utf-8"),
    }


def _prior_message(
    config: dict[str, Any],
    skills: dict[str, str],
    prior: str,
    family: str,
    branch: str,
) -> str:
    if branch == "fallback":
        return (
            "Inherited state\nGLOBAL SLOT: [cold]\nDOMAIN SLOT: [cold]\n"
            "Use only the task evidence and frozen executor contract."
        )
    spec = config["prior_materialization"][prior]
    domain = spec["domain_slot"]
    if domain == "matching_task_family_skill":
        domain = ["officeqa_skill" if family == "OfficeQA" else "spreadsheetbench_skill"]
    def render(names: list[str]) -> str:
        return "\n\n".join(
            f"[{name}]\n{skills[name].strip()}" for name in names
        ) or "[cold]"
    return (
        f"Inherited state type: {prior}\n"
        f"GLOBAL SLOT:\n{render(spec['global_slot'])}\n\n"
        f"DOMAIN SLOT ({family}):\n{render(domain)}"
    )


def _office_messages(
    prior_text: str, row: dict[str, str], evidence: str
) -> list[dict[str, str]]:
    contract = (
        "Use only supplied evidence. Return JSON with exactly answer, evidence, "
        "calculation, abstain, reason. evidence is a list of grounded records. "
        "calculation records operation, time_basis, operands, and result. If the "
        "evidence is insufficient or inconsistent, set abstain=true, use empty "
        "answer/evidence/calculation values, and explain why. No markdown."
    )
    return [
        {"role": "system", "content": contract},
        {"role": "system", "content": prior_text},
        {
            "role": "user",
            "content": f"Question: {row['question']}\n\nEvidence:\n{evidence}",
        },
    ]


def _spreadsheet_messages(
    prior_text: str,
    item: dict[str, Any],
    workbook_path: Path,
    snapshot: dict[str, Any],
) -> list[dict[str, str]]:
    target = str(item["answer_position"])
    contract = (
        "Return JSON with exactly target_range, formula_regions, abstain, reason. "
        f"target_range must be {target}. Each formula_regions item has exactly "
        "range, anchor_cell, formula. A formula is authored at anchor_cell and "
        "will be translated across its range by the deterministic executor. "
        "Regions must exactly cover the target without overlap. If the compact "
        "snapshot is insufficient, abstain instead of guessing. No markdown."
    )
    payload = {
        "instruction": item["instruction"],
        "initial_workbook": relative(workbook_path),
        "target_range": target,
        "workbook_snapshot": snapshot,
    }
    return [
        {"role": "system", "content": contract},
        {"role": "system", "content": prior_text},
        {
            "role": "user",
            "content": json.dumps(payload, ensure_ascii=True, sort_keys=True),
        },
    ]


def build_request_plan(
    config_path: Path = DEFAULT_CONFIG,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    config = validate_config(config_path)
    phase1p = read_json(ROOT / config["source_contracts"]["phase1p_config"]["path"])
    office_rows = _office_rows(
        ROOT / config["source_contracts"]["officeqa_payload"]
    )
    docs_roots = resolve_docs_roots(
        [config["source_contracts"]["officeqa_docs_root"]]
    )
    sheets = {
        str(item["id"]): item
        for item in read_json(
            ROOT / config["source_contracts"]["spreadsheetbench_payload"]
        )
    }
    skills = _skill_texts(config)
    material = config["request_materialization"]
    plan: list[dict[str, Any]] = []
    for prior in material["prior_conditions"]:
        for stratum in material["probe_strata"]:
            for gate in material["gate_policies"]:
                for family in material["task_families"]:
                    ids = phase1p["task_streams"][family][f"{stratum}_probe_ids"]
                    for probe_index, task_id in enumerate(ids, start=1):
                        for branch in material["branches"]:
                            prior_text = _prior_message(
                                config, skills, prior, family, branch
                            )
                            if family == "OfficeQA":
                                row = office_rows[str(task_id)]
                                evidence = compact_office_evidence(
                                    row, docs_roots, material["officeqa_evidence_max_chars"]
                                )
                                messages = _office_messages(prior_text, row, evidence)
                                source = {
                                    "question_sha256": stable_hash(row["question"]),
                                    "evidence_sha256": stable_hash(evidence),
                                    "evidence_chars": len(evidence),
                                }
                            else:
                                item = sheets[str(task_id)]
                                task_dir = (
                                    ROOT
                                    / config["source_contracts"]["spreadsheetbench_root"]
                                    / item["spreadsheet_path"]
                                )
                                initial = sorted(task_dir.glob("*_init.xlsx"))
                                if len(initial) != 1:
                                    raise RunnerError(
                                        f"task {task_id} initial workbook is not unique"
                                    )
                                snapshot = compact_workbook_snapshot(
                                    initial[0],
                                    max_nonempty=material[
                                        "spreadsheet_snapshot_max_nonempty_cells"
                                    ],
                                    head_rows=material[
                                        "spreadsheet_large_sheet_head_rows"
                                    ],
                                    tail_rows=material[
                                        "spreadsheet_large_sheet_tail_rows"
                                    ],
                                )
                                messages = _spreadsheet_messages(
                                    prior_text, item, initial[0], snapshot
                                )
                                source = {
                                    "initial_workbook": relative(initial[0]),
                                    "initial_workbook_sha256": sha256_file(initial[0]),
                                    "snapshot_sha256": stable_hash(snapshot),
                                    "snapshot_nonempty_cells": sum(
                                        len(sheet["cells"])
                                        for sheet in snapshot["sheets"]
                                    ),
                                    "workbook_nonempty_cells": snapshot[
                                        "total_nonempty_cells"
                                    ],
                                    "snapshot_truncated": any(
                                        sheet["snapshot_truncated"]
                                        for sheet in snapshot["sheets"]
                                    ),
                                }
                            request = {
                                "model": config["execution"]["model_id"],
                                "messages": messages,
                                "temperature": config["execution"]["temperature"],
                                "enable_thinking": config["execution"][
                                    "enable_thinking"
                                ],
                            }
                            logical = {
                                "prior_condition": prior,
                                "probe_stratum": stratum,
                                "gate_policy": gate,
                                "task_family": family,
                                "probe_index": probe_index,
                                "task_id": str(task_id),
                                "branch": branch,
                            }
                            logical_call_id = "__".join(
                                str(logical[key])
                                for key in material["logical_call_order"]
                                if key != "probe_index"
                            ) + f"__probe-{probe_index:02d}"
                            plan.append(
                                {
                                    "call_index": len(plan) + 1,
                                    "logical_call_id": logical_call_id,
                                    **logical,
                                    "request": request,
                                    "request_hash": stable_hash(request),
                                    "source": source,
                                }
                            )
    expected = material["logical_calls"]
    if len(plan) != expected or [row["call_index"] for row in plan] != list(
        range(1, expected + 1)
    ):
        raise RunnerError(f"call-count drift: expected {expected}, got {len(plan)}")
    if len({row["logical_call_id"] for row in plan}) != expected:
        raise RunnerError("logical call IDs are not unique")
    if any("max_tokens" in row["request"] for row in plan):
        raise RunnerError("request plan contains forbidden max_tokens")
    return config, plan


def validate_response_contract(family: str, payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise RunnerError("response must be a JSON object")
    required = (
        {"answer", "evidence", "calculation", "abstain", "reason"}
        if family == "OfficeQA"
        else {"target_range", "formula_regions", "abstain", "reason"}
    )
    if set(payload) != required or not isinstance(payload["abstain"], bool):
        raise RunnerError(f"{family} response keys or abstain type changed")
    if payload["abstain"]:
        if not str(payload["reason"]).strip():
            raise RunnerError("abstention requires a nonempty reason")
        return {"decision": "abstain", "contract_valid": True}
    if family == "OfficeQA":
        if not isinstance(payload["evidence"], list) or not payload["evidence"]:
            raise RunnerError("OfficeQA non-abstention requires evidence")
        if not isinstance(payload["calculation"], dict):
            raise RunnerError("OfficeQA calculation must be an object")
    else:
        regions = payload["formula_regions"]
        if not isinstance(regions, list) or not regions:
            raise RunnerError("SpreadsheetBench non-abstention requires formula regions")
        for region in regions:
            if set(region) != {"range", "anchor_cell", "formula"}:
                raise RunnerError("formula region schema changed")
            if not str(region["formula"]).startswith("="):
                raise RunnerError("formula region formula must start with '='")
    return {"decision": "execute", "contract_valid": True}


def exact_usage(response: dict[str, Any]) -> dict[str, int]:
    usage = response.get("usage")
    if not isinstance(usage, dict):
        raise RunnerError("unknown usage")
    try:
        values = {
            "input_tokens": int(usage["input_tokens"]),
            "output_tokens": int(usage["output_tokens"]),
            "total_tokens": int(usage["total_tokens"]),
        }
    except (KeyError, TypeError, ValueError) as exc:
        raise RunnerError("unknown usage") from exc
    if min(values.values()) < 0 or (
        values["input_tokens"] + values["output_tokens"]
        != values["total_tokens"]
    ):
        raise RunnerError("invalid exact usage")
    return values


def _load_records(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [
        json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def execute_plan(
    plan: list[dict[str, Any]],
    output_dir: Path,
    provider: Callable[[dict[str, Any]], dict[str, Any]],
    *,
    max_new_calls: int | None = None,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    plan_hash = stable_hash(plan)
    plan_path = output_dir / "request_plan.json"
    if plan_path.exists():
        if stable_hash(read_json(plan_path)) != plan_hash:
            raise RunnerError("request plan drift")
    else:
        write_json(plan_path, plan)
    results_path = output_dir / "results.jsonl"
    records = _load_records(results_path)
    if len(records) > len(plan):
        raise RunnerError("call-count drift")
    seen: set[str] = set()
    for index, record in enumerate(records):
        expected = plan[index]
        if (
            record.get("logical_call_id") != expected["logical_call_id"]
            or record.get("request_hash") != expected["request_hash"]
            or record.get("call_index") != expected["call_index"]
        ):
            raise RunnerError("completed records are not an exact plan prefix")
        if record["logical_call_id"] in seen:
            raise RunnerError("duplicate logical call ID in resume records")
        seen.add(record["logical_call_id"])
        if record.get("terminal"):
            raise HardStop("terminal failure is not resumable")

    made = 0
    for item in plan[len(records):]:
        if max_new_calls is not None and made >= max_new_calls:
            break
        made += 1
        base = {
            "call_index": item["call_index"],
            "logical_call_id": item["logical_call_id"],
            "request_hash": item["request_hash"],
            "task_family": item["task_family"],
            "task_id": item["task_id"],
            "branch": item["branch"],
            "attempt_count": 1,
        }
        usage: dict[str, int] | None = None
        try:
            response = provider(deepcopy(item["request"]))
            usage = exact_usage(response)
            raw_text = str(response.get("content", ""))
            payload = json.loads(raw_text)
            contract = validate_response_contract(item["task_family"], payload)
            record = {
                **base,
                "status": "completed",
                "terminal": False,
                "usage_known": True,
                **usage,
                "response": payload,
                **contract,
            }
        except Exception as exc:  # noqa: BLE001
            record = {
                **base,
                "status": "hard_stop",
                "terminal": True,
                "usage_known": usage is not None,
                "error_type": type(exc).__name__,
                "error": str(exc),
            }
            if usage is not None:
                record.update(usage)
        with results_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(record, ensure_ascii=True, sort_keys=True) + "\n")
        records.append(record)
        if record["terminal"]:
            break

    completed = sum(record["status"] == "completed" for record in records)
    usage_known = all(record.get("usage_known") is True for record in records)
    decisions = {
        decision: sum(record.get("decision") == decision for record in records)
        for decision in ("execute", "abstain")
    }
    manifest = {
        "schema_version": 1,
        "status": (
            "hard_stopped"
            if records and records[-1].get("terminal")
            else "completed" if completed == len(plan) else "paused"
        ),
        "request_plan_sha256": plan_hash,
        "planned_logical_calls": len(plan),
        "recorded_calls": len(records),
        "completed_calls": completed,
        "provider_attempts": len(records),
        "usage_known_for_all_attempts": usage_known,
        "decisions": decisions,
        "usage": {
            "input_tokens": sum(int(row.get("input_tokens", 0)) for row in records),
            "output_tokens": sum(int(row.get("output_tokens", 0)) for row in records),
            "total_tokens": sum(int(row.get("total_tokens", 0)) for row in records),
        },
    }
    write_json(output_dir / "run_manifest.json", manifest)
    return manifest


def openai_provider(config: dict[str, Any]) -> Callable[[dict[str, Any]], dict[str, Any]]:
    if not config["execution"]["network_calls_allowed"]:
        raise PermissionError("Phase 1R network permission is closed")
    from openai import OpenAI

    key = os.environ.get(config["execution"]["api_key_env"], "")
    if not key.strip():
        raise PermissionError(f"missing {config['execution']['api_key_env']}")
    client = OpenAI(
        api_key=key,
        base_url=config["execution"]["endpoint"].rstrip("/"),
        max_retries=0,
    )
    def call(request: dict[str, Any]) -> dict[str, Any]:
        response = client.chat.completions.create(
            model=request["model"],
            messages=request["messages"],
            temperature=request["temperature"],
            extra_body={"enable_thinking": request["enable_thinking"]},
        )
        usage = getattr(response, "usage", None)
        choices = getattr(response, "choices", None) or []
        return {
            "content": str(getattr(choices[0].message, "content", "")) if choices else "",
            "usage": {
                "input_tokens": getattr(usage, "prompt_tokens", None),
                "output_tokens": getattr(usage, "completion_tokens", None),
                "total_tokens": getattr(usage, "total_tokens", None),
            } if usage is not None else None,
        }
    return call


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--live", action="store_true")
    args = parser.parse_args(argv)
    config, plan = build_request_plan(args.config)
    if not args.live:
        print(json.dumps({
            "mode": "zero-network-plan",
            "logical_calls": len(plan),
            "request_plan_sha256": stable_hash(plan),
            "max_tokens_omitted": all("max_tokens" not in row["request"] for row in plan),
        }, indent=2))
        return 0
    if args.output_dir is None:
        raise RunnerError("--output-dir is required for live execution")
    execute_plan(plan, args.output_dir, openai_provider(config))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
