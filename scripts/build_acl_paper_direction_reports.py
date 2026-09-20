from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "reports"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK = "1F4D78"
GRAY = "F2F4F7"
LIGHT = "F4F6F9"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None: node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    ind = tblPr.find(qn('w:tblInd'))
    if ind is None: ind = OxmlElement('w:tblInd'); tblPr.append(ind)
    ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(w)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i]/1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None: tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'), str(widths[i])); tcW.set(qn('w:type'), 'dxa')

def setup(doc, title, subtitle):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor(0x22,0x22,0x22)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
    for name, size, color, before, after in [('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',12,DARK,8,4)]:
        st=styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(title); r.font.name='Calibri'; r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(DARK)
    p.paragraph_format.space_after=Pt(4)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(subtitle); r.font.name='Calibri'; r.font.size=Pt(11); r.font.italic=True; r.font.color.rgb=RGBColor(0x66,0x66,0x66)
    p.paragraph_format.space_after=Pt(14)
    # quiet footer
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=footer.add_run('ACL 2027 research direction report'); rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(0x77,0x77,0x77)

def p(doc, text, bold_prefix=None):
    para=doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold=True; para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    return para

def bullets(doc, items):
    for item in items:
        para=doc.add_paragraph(style='List Bullet'); para.paragraph_format.space_after=Pt(4); para.add_run(item)

def numbered(doc, items):
    for item in items:
        para=doc.add_paragraph(style='List Number'); para.paragraph_format.space_after=Pt(4); para.add_run(item)

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); set_table_widths(t,widths)
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=h; shade(t.rows[0].cells[i], GRAY)
        for run in t.rows[0].cells[i].paragraphs[0].runs: run.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): cells[i].text=str(val)
    set_table_widths(t,widths)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def build_original():
    d=Document(); setup(d,'原主线论文证据审计报告','ACL 2027 | Historical skill inheritance, typed priors, sparse validation and continual routing')
    p(d,'本报告的目的不是把所有实验包装成正结果，而是逐阶段说明原主线试图证明什么、实际证明了什么，以及哪些主张仍然缺乏证据。')
    d.add_heading('一、原始研究问题与机制链',1)
    p(d,'原主线研究：Agent 能否从历史成功轨迹中抽取 skill，通过 typed/scoped prior 和 sparse validation，在后续任务中提升 continual routing reliability 与 token efficiency？')
    p(d,'机制链：历史成功轨迹 -> skill candidate -> typed/scoped prior -> sparse validation -> accept / reject / abstain -> downstream routing。')
    d.add_heading('二、阶段性实验路线',1)
    d.add_heading('2.1 Phase 0：受控机制验证',2)
    p(d,'Phase 0 在受控或合成环境中操纵 skill identity、scope、prior 类型、probe 代表性和 candidate retention。它支持 identity、scope 和 retention policy 会影响后续行为，并支持 accept/reject/abstain 的操作语义。')
    p(d,'证据边界：Phase 0 证明的是机制可行性，不是机制在真实模型和真实任务上的普遍性能收益。')
    d.add_heading('2.2 Phase 1：真实模型能力底座',2)
    p(d,'Phase 1S 完成 192 次开发调用，仅 6 次精确成功；candidate 分支为 2/96，fresh fallback 为 4/96，representative panel 为 0/96，shifted panel 为 6/96。')
    p(d,'解释：历史经验继承存在能力前置条件。若模型不能稳定地产生可验证成功，就没有可靠 skill 可继承。该结果削弱了乐观假设，但没有单独证明 prior 机制无效。')
    d.add_heading('2.3 Phase 2：真实 typed prior 与 routing',2)
    p(d,'Phase 2 比较 cold、global-only、contextual typed、shuffled 和 incompatible control。coverage incomplete、terminal stop、transport mismatch 和身份不可识别使多个冻结 gate 未通过。')
    p(d,'解释：真实任务中的 prior 效果很容易被模型能力、协议和 payload 可识别性掩盖；没有建立稳定的 typed-prior 因果优势。')
    d.add_heading('2.4 Phase 3：selector uptake 与 answer grounding',2)
    p(d,'Phase 3 将 skill selection 与最终答案分离。Phase 3H 的 contextual target-answer rate 为 0.825，incompatible-control counterfactual-answer rate 为 0.475，paired target-to-counterfactual grounding 为 0.425，冻结 gate 未通过。')
    p(d,'核心发现：selected skill 不等于 executed skill，也不等于 grounded answer。')
    d.add_heading('2.5 Phase 4：协议和传输修复',2)
    p(d,'Phase 4 发现 response contract 没有真正进入模型消息、evidence ID 不可见、部分 global-only/contextual typed payload 完全相同。修复后 development contract gate 通过，但 held-out 仍未产生稳定 typed-prior 优势。')
    d.add_heading('2.6 Phase 5：admission、abstention 与 grounding',2)
    p(d,'Phase 5 完成 240 行，条件为 cold、always-use typed、evidence-grounded admission、shuffled typed、incompatible control 和 evidence abstain。')
    table(d,['条件','答案正确','Evidence resolving','解释'],[
        ['cold','39/40','34/40','cold 并非必然失败'],['always-use typed','40/40','33/40','typed prior 可有效，但 grounding 不稳定'],['admission typed','39/40','33/40','未超过 always-use'],['shuffled typed','32/40','35/40','格式和证据不等于正确执行'],['incompatible control','0/40','35/40','不相容 prior 可污染答案'],['evidence abstain','38/40','33/40','abstention 与答案正确不完全一致']], [1800,1500,1800,4260])
    p(d,'总体：contract-valid 240/240，evidence resolving 203/240，answer correct 188/240，selected skill match 198/240，admission match 145/240。')
    d.add_heading('三、证据审计：已经证明什么',1)
    numbered(d,['历史经验必须带有 identity 和 scope。','不相容 prior 可能系统性污染答案；Phase 5 incompatible control 为 0/40。','skill selection 不等于 skill execution。','contract-valid 不等于 evidence-grounded。','可靠经验继承依赖模型先产生足够可靠的成功。','candidate retention 和 abstention 在受控环境中具有合理语义。','经验复用必须加入 answer-level validation。'])
    d.add_heading('四、证据审计：没有证明什么',1)
    bullets(d,['typed prior 普遍提高真实任务准确率。','evidence-grounded admission 优于 always-use typed。','sparse probe 能代表 downstream distribution。','triage 在真实部署中降低 harmful deployment。','continual routing reliability 或 token efficiency 提升。','跨模型、跨任务、跨领域泛化成立。','历史经验复用已经成为稳定的现实 agent 性能来源。'])
    d.add_heading('五、原主线的结论等级',1)
    p(d,'原命题“Historical experience improves continual agent routing”目前只能算部分支持。更严谨的结论是：受控实验支持 identity、scope 和 retention 机制；真实实验尚未建立稳定收益，但清楚暴露了不相容 prior、selector-grounding gap 和 evidence resolution failure。')
    d.add_heading('六、对原论文的处理建议',1)
    p(d,'原论文不应把所有 protocol readiness 或描述性差异写成方法优越性。若保留原主题，应将主张降级为：历史经验复用的机制条件已被识别，但真实部署收益尚未建立。')
    return d

def build_new():
    d=Document(); setup(d,'新主线论文重构方案','ACL 2027 | Scope-aware admission and answer-level grounding for continual agents')
    p(d,'本报告说明为什么新主线必须从“经验复用是否提升性能”转向“经验何时值得复用，以及为什么 selection 不能保证 answer-level transfer”。')
    d.add_heading('一、新研究问题如何确定',1)
    p(d,'原实验把历史成功、skill 抽象、scope 匹配、admission、操作执行、证据充分性和最终答案压缩成一个 accuracy 指标。结果失败时，无法区分方法失败、模型能力失败、传输失败和 grounding 失败。')
    p(d,'因此新问题改为：Agent 应该在什么条件下复用历史经验？为什么 skill 被选择后仍可能无法形成证据支持的正确答案？')
    d.add_heading('二、新实验的分层框架',1)
    p(d,'新的经验复用链为：verified trajectory -> typed candidate -> scope-aware admission -> evidence-grounded operation -> intermediate result -> final answer。')
    table(d,['层级','需要回答的问题','失败表现'],[
        ['Trajectory','历史成功是否真实可验证？','候选经验本身不可靠'],['Candidate','skill 抽象是否正确？','候选 identity 或 scope 错误'],['Admission','当前任务是否应该使用它？','不该用却接受，或该用却 abstain'],['Operation','模型是否执行了正确操作？','选对 skill 但方向错误'],['Evidence','证据是否支持操作？','ID 不存在、证据不足或错配'],['Answer','最终答案是否被中间过程支持？','格式正确但答案错误']], [1800,3900,3660])
    d.add_heading('三、新实验的核心假设',1)
    d.add_heading('3.1 H1：scope mismatch 会造成系统性错误',2)
    p(d,'不相容 prior 不是普通随机噪声。它应把模型推向可解释的 counterfactual answer。Phase 5 incompatible control 的 0/40 正确率为该假设提供了直接支持。')
    d.add_heading('3.2 H2：evidence-grounded admission 能降低 harmful reuse',2)
    p(d,'这是新方法的核心目标，但当前尚未证明总体准确率优势：always-use typed 为 40/40，admission typed 为 39/40。因此目前只能说 admission 是可审计的安全控制机制，而不能声称已优于 always-use。')
    d.add_heading('3.3 H3：selector uptake 不能代表经验迁移',2)
    p(d,'Phase 3、Phase 3H 和 Phase 5 共同支持 selection != execution != grounding。')
    d.add_heading('3.4 H4：answer grounding 是关键瓶颈',2)
    p(d,'Phase 5 同时得到 240/240 contract-valid、203/240 evidence-resolving 和 188/240 answer-correct，说明格式、证据和答案是三个不同层级。')
    d.add_heading('四、新论文的主要创新点',1)
    numbered(d,['把历史经验形式化为可审计对象：candidate = identity + scope + provenance + evidence + validation status + retention status。','明确区分 selection、execution 和 grounding：选择了 skill 不等于执行了 skill，也不等于得到证据支持的答案。','提出 operation-incompatible control：提供形式上合法但操作不兼容的 prior，用可解释的 counterfactual answer 测量 misuse。','将 answer-level evidence grounding 作为独立评价维度，同时记录 admission、selected skill、evidence IDs、extracted operands、intermediate result 和 final answer。','把负结果转化为安全边界：研究什么条件下经验可安全复用，什么条件下会系统性伤害答案。'])
    d.add_heading('五、推荐的论文中心主张',1)
    p(d,'Historical experience is not safely reusable merely because it is successful, typed, or selected by the agent. Safe reuse requires scope-aware admission and answer-level evidence grounding. Controlled experiments show that incompatible priors can systematically corrupt answers, while real-model experiments reveal that selector uptake and contract validity substantially overestimate true experience transfer.')
    d.add_heading('六、ACL 投稿定位',1)
    p(d,'不建议定位为“我们的方法普遍提升 continual agent performance”。建议定位为：一个用于审计历史经验复用的 scope-aware、evidence-grounded continual-agent framework。')
    d.add_heading('七、建议的论文结构',1)
    numbered(d,['Introduction：提出历史经验复用的安全性问题。','Problem formulation：定义 candidate、scope、admission、evidence 和 grounding。','Failure taxonomy：区分 identity、scope、admission、operation、evidence 和 answer failure。','Controlled mechanism benchmark：用 Phase 0 证明机制在可识别条件下可行。','Real-model evaluation：用 Phase 1-4 展示真实环境中的外部有效性挑战。','Admission and grounding study：用 Phase 5 测量不相容 prior、abstention 和 answer-level grounding。','Implications and limitations：给出适用边界，不夸大跨模型或跨领域泛化。'])
    d.add_heading('八、最终判断',1)
    p(d,'新主线不是放弃原工作，而是把“尚未证实的性能提升主张”重构为“历史经验复用的审计、失败机制和安全边界”。Phase 0 提供机制证据，Phase 1-4 提供真实环境挑战，Phase 5 提供不相容 prior 的直接伤害证据和 selector-grounding gap 的完整测量。')
    return d

if __name__ == '__main__':
    build_original().save(OUT/'ACL2027_原主线论文证据审计报告.docx')
    build_new().save(OUT/'ACL2027_新主线论文重构与创新方案.docx')
    print(OUT/'ACL2027_原主线论文证据审计报告.docx')
    print(OUT/'ACL2027_新主线论文重构与创新方案.docx')
