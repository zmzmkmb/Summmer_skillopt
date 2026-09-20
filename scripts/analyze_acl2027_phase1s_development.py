#!/usr/bin/env python3
"""Score the completed ACL 2027 Phase 1S v2 development-only live run."""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
import tempfile
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import mean
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.run_acl2027_phase1r_development import (
    read_json,
    sha256_file,
    stable_hash,
    write_json,
)

DEFAULT_ARTIFACT = (
    ROOT / "artifacts/acl2027_phase1s_development_method_effect_live_v2"
)
DEFAULT_OFFICEQA = ROOT / "data/officeqa_verified/officeqa_full.csv"
DEFAULT_SPREADSHEET_DATASET = (
    ROOT / "data/spreadsheetbench_verified_400/dataset.json"
)
DEFAULT_PHASE1P_CONFIG = (
    ROOT / "configs/acl2027/phase1p_contribution_aligned_protocol_v1.json"
)

OUTPUT_NAMES = (
    "scored_records.jsonl",
    "analysis_aggregates.json",
    "design_audit.json",
    "analysis_manifest.json",
)


class AnalysisError(RuntimeError):
    """Raised when immutable inputs or scoring invariants fail."""


def _relative(path: Path) -> str:
    try:
        path = path.relative_to(ROOT)
    except ValueError:
        pass
    return str(path).replace("\\", "/")


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def parse_numeric(value: Any) -> Decimal:
    """Parse benchmark numeric answers without changing percent scale."""
    if isinstance(value, bool) or value is None:
        raise AnalysisError(f"non-numeric answer: {value!r}")
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1].strip()
    text = text.replace(",", "").replace("$", "").strip()
    if text.endswith("%"):
        text = text[:-1].strip()
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        raise AnalysisError(f"non-numeric answer: {value!r}")
    try:
        parsed = Decimal(text)
    except InvalidOperation as exc:
        raise AnalysisError(f"non-numeric answer: {value!r}") from exc
    return -parsed if negative else parsed


def _decimal_json(value: Decimal | None) -> int | str | None:
    if value is None:
        return None
    if value == value.to_integral_value():
        return int(value)
    return format(value, "f")


def audit_officeqa_calculation(payload: dict[str, Any]) -> dict[str, Any]:
    calculation = payload.get("calculation")
    if not isinstance(calculation, dict):
        return {"supported": False, "reason": "calculation_not_object"}
    operation = str(calculation.get("operation", "")).strip().lower()
    operands = calculation.get("operands")
    try:
        reported = parse_numeric(calculation.get("result"))
    except AnalysisError:
        return {
            "supported": False,
            "operation": operation,
            "reason": "result_not_numeric",
        }

    recomputed: Decimal | None = None
    if operation in {"sum", "addition", "add"} and isinstance(operands, list):
        try:
            recomputed = sum((parse_numeric(value) for value in operands), Decimal(0))
        except AnalysisError:
            pass
    elif operation == "percent_change":
        try:
            if isinstance(operands, dict):
                initial = parse_numeric(operands["initial_value"])
                final = parse_numeric(operands["final_value"])
            elif isinstance(operands, list) and len(operands) == 2:
                initial = parse_numeric(operands[0])
                final = parse_numeric(operands[1])
            else:
                raise AnalysisError("unsupported percent-change operands")
            if initial != 0:
                recomputed = (final - initial) / initial * Decimal(100)
        except (AnalysisError, KeyError):
            pass
    if recomputed is None:
        return {
            "supported": False,
            "operation": operation,
            "reported_result": _decimal_json(reported),
            "reason": "unsupported_or_unrecomputable",
        }
    return {
        "supported": True,
        "operation": operation,
        "reported_result": _decimal_json(reported),
        "recomputed_result": _decimal_json(recomputed),
        "arithmetic_consistent": abs(recomputed - reported) <= Decimal("0.0005"),
    }


def score_officeqa(
    record: dict[str, Any],
    reference_answer: Any,
) -> dict[str, Any]:
    base = {
        "reference_answer": str(reference_answer),
        "executed": False,
        "exact_task_correct": False,
        "cell_accuracy": 0.0,
        "calculation_audit": {"supported": False, "reason": "not_executed"},
    }
    if record.get("status") != "completed":
        return {**base, "scoring_status": "invalid_output"}
    if record.get("decision") == "abstain":
        return {**base, "scoring_status": "abstain"}
    payload = record.get("response")
    if not isinstance(payload, dict):
        return {**base, "scoring_status": "missing_response"}
    try:
        predicted = parse_numeric(payload.get("answer"))
        reference = parse_numeric(reference_answer)
    except AnalysisError as exc:
        return {
            **base,
            "scoring_status": "non_numeric_answer",
            "answer_error": str(exc),
            "calculation_audit": audit_officeqa_calculation(payload),
        }
    correct = predicted == reference
    return {
        **base,
        "scoring_status": "executed",
        "executed": True,
        "predicted_answer": _decimal_json(predicted),
        "normalized_reference_answer": _decimal_json(reference),
        "exact_task_correct": correct,
        "cell_accuracy": float(correct),
        "calculation_audit": audit_officeqa_calculation(payload),
    }


def split_qualified_reference(
    reference: str,
    *,
    default_sheet: str | None = None,
) -> tuple[str | None, str]:
    text = str(reference).strip()
    if "!" not in text:
        return default_sheet, text.replace("$", "")
    sheet, coordinate = text.rsplit("!", 1)
    sheet = sheet.strip()
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, coordinate.strip().replace("$", "")


def _range_cells(reference: str) -> list[str]:
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    return [
        f"{get_column_letter(column)}{row}"
        for row in range(min_row, max_row + 1)
        for column in range(min_col, max_col + 1)
    ]


def expand_formula_regions(
    response: dict[str, Any],
    *,
    expected_target_range: str,
    default_sheet: str,
) -> tuple[str, dict[str, str]]:
    response_target = str(response.get("target_range", "")).strip()
    expected_sheet, expected_range = split_qualified_reference(
        expected_target_range, default_sheet=default_sheet
    )
    observed_sheet, observed_range = split_qualified_reference(
        response_target, default_sheet=default_sheet
    )
    if observed_sheet != expected_sheet or observed_range.upper() != expected_range.upper():
        raise AnalysisError("response target range does not match benchmark target")
    regions = response.get("formula_regions")
    if not isinstance(regions, list) or not regions:
        raise AnalysisError("formula regions are missing")

    expected_cells = set(_range_cells(expected_range))
    formulas: dict[str, str] = {}
    for region in regions:
        if not isinstance(region, dict):
            raise AnalysisError("formula region is not an object")
        region_sheet, region_range = split_qualified_reference(
            str(region.get("range", "")), default_sheet=expected_sheet
        )
        anchor_sheet, anchor = split_qualified_reference(
            str(region.get("anchor_cell", "")), default_sheet=region_sheet
        )
        if region_sheet != expected_sheet or anchor_sheet != expected_sheet:
            raise AnalysisError("formula region crosses the benchmark target sheet")
        region_cells = _range_cells(region_range)
        min_col, min_row, max_col, max_row = range_boundaries(region_range)
        anchor_row, anchor_col = coordinate_to_tuple(anchor)
        if not (min_row <= anchor_row <= max_row and min_col <= anchor_col <= max_col):
            raise AnalysisError("formula anchor is outside its region")
        formula = str(region.get("formula", ""))
        if not formula.startswith("="):
            raise AnalysisError("formula does not start with '='")
        for cell in region_cells:
            if cell in formulas:
                raise AnalysisError(f"overlapping formula regions at {cell}")
            try:
                formulas[cell] = Translator(formula, origin=anchor).translate_formula(cell)
            except Exception as exc:  # noqa: BLE001
                raise AnalysisError(f"formula translation failed at {cell}: {exc}") from exc
    if set(formulas) != expected_cells:
        missing = sorted(expected_cells - set(formulas))
        extra = sorted(set(formulas) - expected_cells)
        raise AnalysisError(
            f"formula regions do not exactly cover target; missing={missing}, extra={extra}"
        )
    return str(expected_sheet), formulas


def _formula_text(value: Any) -> Any:
    text = getattr(value, "text", None)
    return text if text is not None else value


def score_spreadsheet(
    record: dict[str, Any],
    *,
    expected_target_range: str,
    initial_workbook: Path,
    golden_workbook: Path,
    temp_root: Path,
    score_cache: dict[str, dict[str, Any]] | None = None,
) -> dict[str, Any]:
    workbook = load_workbook(golden_workbook, data_only=False, read_only=True)
    try:
        target_sheet, bare_target = split_qualified_reference(
            expected_target_range,
            default_sheet=workbook.sheetnames[0],
        )
        if target_sheet not in workbook.sheetnames:
            raise AnalysisError(f"golden target sheet not found: {target_sheet}")
        target_count = len(_range_cells(bare_target))
    finally:
        workbook.close()
    base = {
        "benchmark_target_range": expected_target_range,
        "target_cells": target_count,
        "executed": False,
        "exact_task_correct": False,
        "cell_accuracy": 0.0,
        "formula_matches": 0,
        "formula_mismatches": [],
    }
    if record.get("status") != "completed":
        return {**base, "scoring_status": "invalid_output"}
    if record.get("decision") == "abstain":
        return {**base, "scoring_status": "abstain"}
    payload = record.get("response")
    if not isinstance(payload, dict):
        return {**base, "scoring_status": "missing_response"}
    try:
        sheet_name, formulas = expand_formula_regions(
            payload,
            expected_target_range=expected_target_range,
            default_sheet=target_sheet,
        )
    except AnalysisError as exc:
        return {
            **base,
            "scoring_status": "executor_rejected",
            "executor_error": str(exc),
        }

    cache_key = stable_hash(
        {
            "expected_target_range": expected_target_range,
            "initial_workbook_sha256": sha256_file(initial_workbook),
            "golden_workbook_sha256": sha256_file(golden_workbook),
            "sheet_name": sheet_name,
            "formulas": formulas,
        }
    )
    if score_cache is not None and cache_key in score_cache:
        return {**base, **score_cache[cache_key], "persisted_score_cache_hit": True}

    with tempfile.TemporaryDirectory(prefix="phase1s_score_", dir=temp_root) as tmp:
        generated = Path(tmp) / "generated.xlsx"
        shutil.copy2(initial_workbook, generated)
        predicted = load_workbook(generated, data_only=False)
        try:
            if sheet_name not in predicted.sheetnames:
                raise AnalysisError(f"initial target sheet not found: {sheet_name}")
            for cell, formula in formulas.items():
                predicted[sheet_name][cell] = formula
            predicted.save(generated)
        finally:
            predicted.close()

        persisted = load_workbook(generated, data_only=False, read_only=True)
        golden = load_workbook(golden_workbook, data_only=False, read_only=True)
        try:
            mismatches = []
            matches = 0
            for cell in sorted(formulas, key=coordinate_to_tuple):
                actual = _formula_text(persisted[sheet_name][cell].value)
                expected = _formula_text(golden[sheet_name][cell].value)
                if actual == expected:
                    matches += 1
                else:
                    mismatches.append(
                        {"cell": cell, "predicted": actual, "golden": expected}
                    )
        finally:
            persisted.close()
            golden.close()
    exact = matches == target_count
    executed_score = {
        "scoring_status": "executed",
        "executed": True,
        "formula_matches": matches,
        "formula_mismatches": mismatches,
        "exact_task_correct": exact,
        "cell_accuracy": matches / target_count,
        "persisted_score_cache_hit": False,
    }
    if score_cache is not None:
        score_cache[cache_key] = executed_score
    return {**base, **executed_score}


def _summarize(rows: Iterable[dict[str, Any]]) -> dict[str, Any]:
    values = list(rows)
    n = len(values)
    if not n:
        return {"calls": 0}
    return {
        "calls": n,
        "contract_valid": sum(bool(row["contract_valid"]) for row in values),
        "invalid_output": sum(row["status"] == "invalid_output" for row in values),
        "execute": sum(row.get("decision") == "execute" for row in values),
        "abstain": sum(row.get("decision") == "abstain" for row in values),
        "exact_task_correct": sum(bool(row["exact_task_correct"]) for row in values),
        "exact_task_accuracy": mean(bool(row["exact_task_correct"]) for row in values),
        "mean_cell_accuracy": mean(float(row["cell_accuracy"]) for row in values),
        "usage": {
            key: sum(int(row[key]) for row in values)
            for key in ("input_tokens", "output_tokens", "total_tokens")
        },
    }


def _group_summaries(
    rows: list[dict[str, Any]],
    dimensions: tuple[str, ...],
) -> list[dict[str, Any]]:
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[tuple(row[key] for key in dimensions)].append(row)
    return [
        {
            **dict(zip(dimensions, key)),
            **_summarize(group),
        }
        for key, group in sorted(groups.items(), key=lambda item: item[0])
    ]


def paired_margins(rows: list[dict[str, Any]]) -> dict[str, Any]:
    pair_dimensions = (
        "prior_condition",
        "probe_stratum",
        "gate_policy",
        "task_family",
        "task_id",
        "probe_index",
    )
    pairs: dict[tuple[Any, ...], dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in rows:
        pairs[tuple(row[key] for key in pair_dimensions)][row["branch"]] = row
    if any(set(branches) != {"candidate", "fallback"} for branches in pairs.values()):
        raise AnalysisError("candidate/fallback pairing is incomplete")
    pair_rows = []
    for key, branches in sorted(pairs.items()):
        candidate = branches["candidate"]
        fallback = branches["fallback"]
        pair_rows.append(
            {
                **dict(zip(pair_dimensions, key)),
                "candidate_correct": bool(candidate["exact_task_correct"]),
                "fallback_correct": bool(fallback["exact_task_correct"]),
                "exact_task_margin": int(bool(candidate["exact_task_correct"]))
                - int(bool(fallback["exact_task_correct"])),
                "cell_accuracy_margin": float(candidate["cell_accuracy"])
                - float(fallback["cell_accuracy"]),
                "candidate_total_tokens": int(candidate["total_tokens"]),
                "fallback_total_tokens": int(fallback["total_tokens"]),
            }
        )

    def aggregate(dimensions: tuple[str, ...]) -> list[dict[str, Any]]:
        groups: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
        for row in pair_rows:
            groups[tuple(row[key] for key in dimensions)].append(row)
        return [
            {
                **dict(zip(dimensions, key)),
                "pairs": len(group),
                "mean_exact_task_margin": mean(
                    row["exact_task_margin"] for row in group
                ),
                "mean_cell_accuracy_margin": mean(
                    row["cell_accuracy_margin"] for row in group
                ),
                "candidate_wins": sum(row["exact_task_margin"] > 0 for row in group),
                "ties": sum(row["exact_task_margin"] == 0 for row in group),
                "fallback_wins": sum(row["exact_task_margin"] < 0 for row in group),
            }
            for key, group in sorted(groups.items(), key=lambda item: item[0])
        ]

    condition_dimensions = (
        "prior_condition",
        "probe_stratum",
        "gate_policy",
        "task_family",
    )
    return {
        "pair_count": len(pair_rows),
        "pairs": pair_rows,
        "by_condition": aggregate(condition_dimensions),
        "by_prior_condition": aggregate(("prior_condition",)),
        "by_probe_stratum": aggregate(("probe_stratum",)),
        "by_gate_policy": aggregate(("gate_policy",)),
        "by_task_family": aggregate(("task_family",)),
    }


def build_design_audit(
    plan: list[dict[str, Any]],
    phase1p: dict[str, Any],
) -> dict[str, Any]:
    gate_pair_dimensions = (
        "prior_condition",
        "probe_stratum",
        "task_family",
        "task_id",
        "probe_index",
        "branch",
    )
    gate_pairs: dict[tuple[Any, ...], dict[str, str]] = defaultdict(dict)
    for row in plan:
        gate_pairs[tuple(row[key] for key in gate_pair_dimensions)][
            row["gate_policy"]
        ] = row["request_hash"]
    complete_gate_pairs = [
        hashes for hashes in gate_pairs.values()
        if set(hashes) == {"candidate_retaining_triage", "destructive_gate"}
    ]
    identical_gate_pairs = sum(
        len(set(hashes.values())) == 1 for hashes in complete_gate_pairs
    )
    request_counts = Counter(row["request_hash"] for row in plan)
    multiplicities = Counter(request_counts.values())

    downstream_by_family = {}
    for family, stream in phase1p["task_streams"].items():
        observed = {row["task_id"] for row in plan if row["task_family"] == family}
        held_out = {str(task_id) for task_id in stream["held_out_ids"]}
        downstream_by_family[family] = {
            "observed_task_ids": sorted(observed),
            "held_out_task_ids": sorted(held_out),
            "observed_held_out_intersection": sorted(observed & held_out),
        }
    no_downstream = all(
        not item["observed_held_out_intersection"]
        for item in downstream_by_family.values()
    )
    gate_labels_only = (
        len(complete_gate_pairs) == len(gate_pairs)
        and identical_gate_pairs == len(complete_gate_pairs)
    )
    return {
        "planned_calls": len(plan),
        "unique_request_bodies": len(request_counts),
        "request_hash_multiplicity_histogram": {
            str(key): value for key, value in sorted(multiplicities.items())
        },
        "gate_label_pairs": len(complete_gate_pairs),
        "gate_label_pairs_with_identical_requests": identical_gate_pairs,
        "gate_policies_implement_distinct_request_or_decision_rules": False,
        "gate_labels_are_metadata_only_in_this_run": gate_labels_only,
        "numerical_gate_thresholds_implemented": False,
        "downstream_tasks_executed": not no_downstream,
        "downstream_task_audit": downstream_by_family,
        "claims_identified": {
            "development_candidate_vs_fallback_effect": True,
            "development_prior_condition_difference": True,
            "development_representative_vs_shifted_difference": True,
            "probe_to_downstream_representativeness": False,
            "false_safe_or_false_harm_deployment_rates": False,
            "triage_superiority_over_destructive_gating": False,
            "candidate_retention_benefit_over_time": False,
        },
        "required_interpretation": (
            "Development-only branch and condition effects are measurable. "
            "Downstream representativeness and gate-policy efficacy are not "
            "identified because no held-out stream was executed and both gate "
            "labels reuse identical requests without numerical decision rules."
        ),
    }


def analyze(
    *,
    artifact_dir: Path = DEFAULT_ARTIFACT,
    officeqa_path: Path = DEFAULT_OFFICEQA,
    spreadsheet_dataset_path: Path = DEFAULT_SPREADSHEET_DATASET,
    phase1p_config_path: Path = DEFAULT_PHASE1P_CONFIG,
) -> dict[str, Any]:
    artifact_dir = artifact_dir.resolve()
    for name in OUTPUT_NAMES:
        if (artifact_dir / name).exists():
            raise AnalysisError(f"refusing to overwrite analysis output: {name}")
    plan_path = artifact_dir / "request_plan.json"
    results_path = artifact_dir / "results.jsonl"
    manifest_path = artifact_dir / "run_manifest.json"
    plan = read_json(plan_path)
    records = _load_jsonl(results_path)
    live_manifest = read_json(manifest_path)
    if len(plan) != 192 or len(records) != 192:
        raise AnalysisError("Phase 1S v2 artifact is not the complete 192-call run")
    if live_manifest.get("provider_attempts") != 192:
        raise AnalysisError("live manifest provider-attempt count drift")
    for item, record in zip(plan, records):
        if any(
            record.get(key) != item.get(key)
            for key in ("call_index", "logical_call_id", "request_hash", "task_id")
        ):
            raise AnalysisError("results are not an exact request-plan prefix")

    with officeqa_path.open(encoding="utf-8-sig", newline="") as handle:
        officeqa = {str(row["uid"]): row for row in csv.DictReader(handle)}
    spreadsheet_items = {
        str(item["id"]): item for item in read_json(spreadsheet_dataset_path)
    }
    phase1p = read_json(phase1p_config_path)

    scored: list[dict[str, Any]] = []
    spreadsheet_score_cache: dict[str, dict[str, Any]] = {}
    with tempfile.TemporaryDirectory(prefix="phase1s_analysis_") as tmp:
        temp_root = Path(tmp)
        for item, record in zip(plan, records):
            dimensions = {
                key: item[key]
                for key in (
                    "prior_condition",
                    "probe_stratum",
                    "gate_policy",
                    "task_family",
                    "task_id",
                    "probe_index",
                    "branch",
                )
            }
            if item["task_family"] == "OfficeQA":
                if item["task_id"] not in officeqa:
                    raise AnalysisError(f"OfficeQA reference missing: {item['task_id']}")
                score = score_officeqa(
                    record, officeqa[item["task_id"]]["answer"]
                )
            else:
                task = spreadsheet_items.get(item["task_id"])
                if task is None:
                    raise AnalysisError(
                        f"SpreadsheetBench task missing: {item['task_id']}"
                    )
                initial = ROOT / item["source"]["initial_workbook"]
                golden = sorted(initial.parent.glob("*_golden.xlsx"))
                if len(golden) != 1:
                    raise AnalysisError(
                        f"golden workbook is not unique: {item['task_id']}"
                    )
                score = score_spreadsheet(
                    record,
                    expected_target_range=str(task["answer_position"]),
                    initial_workbook=initial,
                    golden_workbook=golden[0],
                    temp_root=temp_root,
                    score_cache=spreadsheet_score_cache,
                )
            scored.append(
                {
                    "call_index": record["call_index"],
                    "logical_call_id": record["logical_call_id"],
                    "request_hash": record["request_hash"],
                    **dimensions,
                    "status": record["status"],
                    "contract_valid": bool(record.get("contract_valid")),
                    "decision": record.get("decision", "invalid"),
                    "input_tokens": int(record["input_tokens"]),
                    "output_tokens": int(record["output_tokens"]),
                    "total_tokens": int(record["total_tokens"]),
                    **score,
                }
            )

    scored_path = artifact_dir / "scored_records.jsonl"
    with scored_path.open("x", encoding="utf-8") as handle:
        for row in scored:
            handle.write(json.dumps(row, ensure_ascii=True, sort_keys=True) + "\n")

    condition_dimensions = (
        "prior_condition",
        "probe_stratum",
        "gate_policy",
        "task_family",
        "branch",
    )
    aggregates = {
        "overall": _summarize(scored),
        "condition_cells": _group_summaries(scored, condition_dimensions),
        "by_prior_condition": _group_summaries(scored, ("prior_condition",)),
        "by_probe_stratum": _group_summaries(scored, ("probe_stratum",)),
        "by_gate_policy": _group_summaries(scored, ("gate_policy",)),
        "by_task_family": _group_summaries(scored, ("task_family",)),
        "by_branch": _group_summaries(scored, ("branch",)),
        "by_task": _group_summaries(scored, ("task_family", "task_id")),
        "paired_margins": paired_margins(scored),
    }
    aggregate_path = artifact_dir / "analysis_aggregates.json"
    write_json(aggregate_path, aggregates)

    design_audit = build_design_audit(plan, phase1p)
    design_path = artifact_dir / "design_audit.json"
    write_json(design_path, design_audit)

    output_hashes = {
        "scored_records": sha256_file(scored_path),
        "analysis_aggregates": sha256_file(aggregate_path),
        "design_audit": sha256_file(design_path),
    }
    analysis_manifest = {
        "schema_version": 1,
        "analysis": "acl2027_phase1s_development_method_effect_live_v2",
        "status": "completed",
        "network_calls": 0,
        "provider_calls": 0,
        "paid_api_calls": 0,
        "input_artifact": _relative(artifact_dir),
        "immutable_input_hashes": {
            "request_plan": sha256_file(plan_path),
            "results": sha256_file(results_path),
            "live_run_manifest": sha256_file(manifest_path),
            "officeqa_reference": sha256_file(officeqa_path),
            "spreadsheetbench_dataset": sha256_file(spreadsheet_dataset_path),
            "phase1p_config": sha256_file(phase1p_config_path),
        },
        "scored_calls": len(scored),
        "output_hashes": output_hashes,
        "aggregate_fingerprint": stable_hash(
            {"outputs": output_hashes, "overall": aggregates["overall"]}
        ),
        "scientific_scope": design_audit["required_interpretation"],
    }
    write_json(artifact_dir / "analysis_manifest.json", analysis_manifest)
    return analysis_manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifact-dir", type=Path, default=DEFAULT_ARTIFACT)
    parser.add_argument("--officeqa", type=Path, default=DEFAULT_OFFICEQA)
    parser.add_argument(
        "--spreadsheet-dataset",
        type=Path,
        default=DEFAULT_SPREADSHEET_DATASET,
    )
    parser.add_argument(
        "--phase1p-config",
        type=Path,
        default=DEFAULT_PHASE1P_CONFIG,
    )
    args = parser.parse_args(argv)
    manifest = analyze(
        artifact_dir=args.artifact_dir,
        officeqa_path=args.officeqa,
        spreadsheet_dataset_path=args.spreadsheet_dataset,
        phase1p_config_path=args.phase1p_config,
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
