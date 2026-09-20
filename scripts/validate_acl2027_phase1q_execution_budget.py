#!/usr/bin/env python3
"""Audit the zero-network ACL 2027 Phase 1Q execution and budget plan."""
from __future__ import annotations

import csv
import hashlib
import json
import sys
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from skillopt.envs.officeqa.tool_runtime import (
    _iter_oracle_refs,
    _locate_parsed_json,
    _render_parsed_page,
    resolve_docs_roots,
)

CONFIG_PATH = ROOT / "configs/acl2027/phase1q_execution_budget_preflight_v1.json"
OUTPUT_DIR = ROOT / "artifacts/acl2027_phase1q_execution_budget_preflight_v1"


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _relative(path: Path) -> str:
    return str(path.relative_to(ROOT)).replace("\\", "/")


def _cold_or_execute(eligible: bool, eligible_route: str) -> str:
    return eligible_route if eligible else "abstain"


def _office_rows(path: Path) -> dict[str, dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for raw in csv.DictReader(handle):
            uid = str(raw.get("uid", "")).strip()
            if not uid:
                continue
            rows[uid] = {
                "uid": uid,
                "question": str(raw.get("question", "")).strip(),
                "source_docs": str(raw.get("source_docs", "")).strip(),
                "source_files": str(raw.get("source_files", "")).strip(),
                "difficulty": str(raw.get("difficulty", "")).strip(),
            }
    return rows


def _task_ids(phase1p: dict[str, Any], family: str) -> list[str]:
    stream = phase1p["task_streams"][family]
    return [
        str(item)
        for item in stream["development_ids"] + stream["held_out_ids"]
    ]


def audit_officeqa(
    phase1p: dict[str, Any],
    config: dict[str, Any],
) -> list[dict[str, Any]]:
    source = config["source_contracts"]
    rows = _office_rows(ROOT / source["officeqa_payload"])
    docs_roots = resolve_docs_roots([source["officeqa_docs_root"]])
    contract = config["executor_eligibility"]["OfficeQA"]
    records: list[dict[str, Any]] = []
    for task_id in _task_ids(phase1p, "OfficeQA"):
        row = rows.get(task_id)
        refs = (
            _iter_oracle_refs(row["source_files"], row["source_docs"])
            if row
            else []
        )
        pages: list[dict[str, Any]] = []
        for source_file, page_number, _ in refs:
            parsed_path = _locate_parsed_json(source_file, docs_roots)
            rendered = (
                _render_parsed_page(str(parsed_path), page_number)
                if parsed_path is not None
                else ""
            )
            pages.append(
                {
                    "source_file": source_file,
                    "page_number": page_number,
                    "parsed_json_exists": parsed_path is not None,
                    "rendered_chars": len(rendered),
                }
            )
        checks = {
            "payload_row_exists": row is not None,
            "question_present": bool(row and row["question"]),
            "source_reference_present": bool(refs),
            "parsed_source_json_exists": bool(pages)
            and all(item["parsed_json_exists"] for item in pages),
            "referenced_page_has_local_content": bool(pages)
            and all(item["rendered_chars"] > 0 for item in pages),
        }
        eligible = all(checks.values())
        records.append(
            {
                "task_family": "OfficeQA",
                "task_id": task_id,
                "eligible": eligible,
                "decision": _cold_or_execute(
                    eligible,
                    contract["eligible_route"],
                ),
                "checks": checks,
                "source_page_count": len(pages),
                "source_pages": pages,
                "reference_answer_accessed": False,
            }
        )
    return records


def _answer_location(
    answer_position: str,
    workbook_sheets: list[str],
) -> tuple[str, str]:
    text = str(answer_position).strip()
    if "!" in text:
        sheet_text, cell_range = text.rsplit("!", 1)
        sheet = sheet_text.strip().strip("'\"")
    else:
        sheet = workbook_sheets[0] if workbook_sheets else ""
        cell_range = text
    return sheet.rstrip("'").lstrip("'"), cell_range.strip().strip("'\"")


def audit_spreadsheetbench(
    phase1p: dict[str, Any],
    config: dict[str, Any],
) -> list[dict[str, Any]]:
    source = config["source_contracts"]
    payload = {
        str(item["id"]): item
        for item in read_json(ROOT / source["spreadsheetbench_payload"])
    }
    contract = config["executor_eligibility"]["SpreadsheetBench"]
    records: list[dict[str, Any]] = []
    for task_id in _task_ids(phase1p, "SpreadsheetBench"):
        item = payload.get(task_id)
        task_dir = (
            ROOT / source["spreadsheetbench_root"] / item["spreadsheet_path"]
            if item
            else ROOT / "__missing_phase1q_task__"
        )
        initial = sorted(task_dir.glob("*_init.xlsx"))
        golden = sorted(task_dir.glob("*_golden.xlsx"))
        workbook_open = False
        initial_sheets: list[str] = []
        golden_sheets: list[str] = []
        try:
            if len(initial) == 1 and len(golden) == 1:
                initial_book = load_workbook(
                    initial[0],
                    read_only=True,
                    data_only=False,
                )
                golden_book = load_workbook(
                    golden[0],
                    read_only=True,
                    data_only=False,
                )
                initial_sheets = list(initial_book.sheetnames)
                golden_sheets = list(golden_book.sheetnames)
                initial_book.close()
                golden_book.close()
                workbook_open = True
        except Exception:  # noqa: BLE001
            workbook_open = False

        target_sheet = ""
        target_range = ""
        range_parseable = False
        target_cell_count = 0
        try:
            target_sheet, target_range = _answer_location(
                item["answer_position"],
                initial_sheets,
            )
            min_col, min_row, max_col, max_row = range_boundaries(target_range)
            range_parseable = (
                min_col >= 1
                and min_row >= 1
                and max_col >= min_col
                and max_row >= min_row
            )
            if range_parseable:
                target_cell_count = (
                    (max_col - min_col + 1) * (max_row - min_row + 1)
                )
        except (KeyError, TypeError, ValueError):
            range_parseable = False

        checks = {
            "payload_row_exists": item is not None,
            "cell_level_manipulation": bool(
                item
                and item.get("instruction_type") == "Cell-Level Manipulation"
            ),
            "one_initial_workbook": len(initial) == 1,
            "one_golden_workbook": len(golden) == 1,
            "answer_range_parseable": range_parseable,
            "target_sheet_exists": bool(target_sheet)
            and target_sheet in initial_sheets
            and target_sheet in golden_sheets,
            "both_workbooks_open": workbook_open,
        }
        eligible = all(checks.values())
        records.append(
            {
                "task_family": "SpreadsheetBench",
                "task_id": task_id,
                "eligible": eligible,
                "decision": _cold_or_execute(
                    eligible,
                    contract["eligible_route"],
                ),
                "checks": checks,
                "target_sheet": target_sheet,
                "target_range": target_range,
                "target_cell_count": target_cell_count,
                "initial_workbook": _relative(initial[0])
                if len(initial) == 1
                else None,
                "golden_workbook": _relative(golden[0])
                if len(golden) == 1
                else None,
                "golden_cell_values_accessed": False,
            }
        )
    return records


def build_call_matrix(config: dict[str, Any]) -> list[dict[str, Any]]:
    accounting = config["call_accounting"]
    cost = config["cost_planning"]
    cells_per_family = (
        accounting["prior_conditions"]
        * accounting["probe_strata"]
        * accounting["gate_policies"]
    )
    branches = accounting["counterfactual_branches_per_task"]
    rates = cost["planning_ceiling_rates_cny_per_million_tokens"]
    rows: list[dict[str, Any]] = []
    for plan_name, plan in accounting["plans"].items():
        tasks_per_cell = (
            accounting["probe_tasks_per_cell"]
            + plan["held_out_tasks_per_cell"]
        )
        family_calls: dict[str, int] = {}
        family_tokens: dict[str, dict[str, int]] = {}
        for family, envelope in cost["per_call_planning_envelopes"].items():
            calls = cells_per_family * tasks_per_cell * branches
            family_calls[family] = calls
            family_tokens[family] = {
                "input_tokens": calls * envelope["input_tokens"],
                "output_tokens": calls * envelope["output_tokens"],
            }
        input_tokens = sum(
            item["input_tokens"] for item in family_tokens.values()
        )
        output_tokens = sum(
            item["output_tokens"] for item in family_tokens.values()
        )
        total_calls = sum(family_calls.values())
        cny = (
            input_tokens * rates["input"]
            + output_tokens * rates["output"]
        ) / 1_000_000
        rows.append(
            {
                "plan": plan_name,
                "tasks_per_cell": tasks_per_cell,
                "held_out_tasks_per_cell": plan["held_out_tasks_per_cell"],
                "condition_accounted_calls": total_calls,
                "physical_request_ceiling": total_calls,
                "candidate_calls": total_calls // branches,
                "fallback_calls": total_calls // branches,
                "max_provider_attempts": total_calls,
                "family_calls": family_calls,
                "planning_input_tokens": input_tokens,
                "planning_output_tokens": output_tokens,
                "planning_total_tokens": input_tokens + output_tokens,
                "planning_ceiling_cny": round(cny, 6),
                "provider_quote": False,
                "request_max_tokens_omitted": True,
            }
        )
    return rows


def audit(config: dict[str, Any]) -> dict[str, Any]:
    source = config["source_contracts"]
    hashes = {
        name: sha256_file(ROOT / item["path"])
        for name, item in source.items()
        if isinstance(item, dict) and "path" in item
    }
    phase1p = read_json(ROOT / source["phase1p_config"]["path"])
    phase1p_manifest = read_json(ROOT / source["phase1p_manifest"]["path"])
    office = audit_officeqa(phase1p, config)
    spreadsheet = audit_spreadsheetbench(phase1p, config)
    eligibility = office + spreadsheet
    matrix = build_call_matrix(config)

    bindings = config["candidate_provenance"]["condition_bindings"]
    binding_pairs = {
        (item["prior_condition"], item["gate_policy"]) for item in bindings
    }
    expected_pairs = {
        (prior, gate)
        for prior in phase1p["prior_conditions"]
        if prior != "identity_invariant"
        for gate in phase1p["gate_policies"]
    }
    staged = config["call_accounting"]["plans"]["staged_confirmation"][
        "held_out_prefix_ids"
    ]
    execution = config["execution"]
    cost = config["cost_planning"]
    checks = {
        "source_hashes_match": all(
            hashes[name] == source[name]["sha256"] for name in hashes
        ),
        "phase1p_is_complete_and_immutable": (
            phase1p_manifest["complete_grid"]
            and phase1p_manifest["aggregate_fingerprint"]
            == "8f5f87370999567c40e86ad40bac4edca12c15a60261384dd8952c7b40013968"
        ),
        "zero_network_paid_and_scaling": (
            execution["protocol_only"]
            and not execution["paid_api_allowed"]
            and not execution["network_calls_allowed"]
            and execution["provider_calls"] == 0
            and not execution["formal_scaling_allowed"]
            and not execution["qwen3_8_max_allowed"]
            and not execution["legacy_officeqa_24_batch_allowed"]
            and not execution["legacy_spreadsheetbench_40_batch_allowed"]
        ),
        "max_tokens_omitted_contract": (
            "max_tokens" in execution["request_body_forbidden_keys"]
            and cost["request_max_tokens_omitted"]
        ),
        "zero_retry_exact_attempt_contract": (
            execution["max_provider_attempts_per_logical_call"] == 1
            and not execution["retries_allowed"]
        ),
        "skill_hashes_match": all(
            hashes[name] == source[name]["sha256"]
            for name in ("officeqa_skill", "spreadsheetbench_skill")
        ),
        "candidate_bindings_cover_all_prior_gate_pairs": (
            len(bindings) == 6
            and len({item["candidate_id"] for item in bindings}) == 6
            and binding_pairs == expected_pairs
        ),
        "fallback_is_fresh_cold_and_shared_by_identity_only": (
            all(
                item["fallback_id"] == "fresh_cold_executor_backed_v1"
                for item in bindings
            )
            and not config["candidate_provenance"]["fallback"][
                "state_reused_across_conditions"
            ]
            and config["candidate_provenance"]["fallback"][
                "candidate_skill_content_omitted"
            ]
        ),
        "prior_bindings_are_distinct": len(
            {
                json.dumps(value, sort_keys=True)
                for value in config["candidate_provenance"][
                    "prior_bindings"
                ].values()
            }
        )
        == 3,
        "all_48_identities_audited": (
            len(eligibility) == 48
            and len(
                {
                    (item["task_family"], item["task_id"])
                    for item in eligibility
                }
            )
            == 48
        ),
        "all_eligibility_decisions_explicit": all(
            item["decision"] != "" for item in eligibility
        ),
        "unsupported_tasks_route_to_abstain": all(
            item["eligible"] or item["decision"] == "abstain"
            for item in eligibility
        )
        and not config["executor_eligibility"]["silent_exclusion_allowed"],
        "officeqa_reference_answers_not_accessed": all(
            not item["reference_answer_accessed"] for item in office
        ),
        "spreadsheet_golden_values_not_accessed": all(
            not item["golden_cell_values_accessed"] for item in spreadsheet
        ),
        "staged_prefix_is_first_four_stream_tasks": all(
            staged[family]
            == phase1p["task_streams"][family]["stream_order"][:4]
            for family in ("OfficeQA", "SpreadsheetBench")
        ),
        "exact_call_counts_are_192_384_960": {
            row["plan"]: row["condition_accounted_calls"] for row in matrix
        }
        == {
            "development_only": 192,
            "staged_confirmation": 384,
            "full_frozen_design": 960,
        },
        "candidate_and_fallback_calls_are_equal": all(
            row["candidate_calls"] == row["fallback_calls"]
            for row in matrix
        ),
        "physical_requests_not_undercounted": all(
            row["physical_request_ceiling"]
            == row["condition_accounted_calls"]
            == row["max_provider_attempts"]
            for row in matrix
        ),
        "planning_rates_are_not_provider_quotes": (
            not any(row["provider_quote"] for row in matrix)
            and cost["provider_price_verification_required_before_paid_authorization"]
            and not cost["hard_execution_budget_authorized"]
        ),
    }
    eligible_count = sum(item["eligible"] for item in eligibility)
    return {
        "analysis": "phase1q_zero_network_execution_budget_preflight",
        "source_hashes": hashes,
        "candidate_provenance": deepcopy(config["candidate_provenance"]),
        "eligibility_summary": {
            "audited": len(eligibility),
            "eligible": eligible_count,
            "abstain": len(eligibility) - eligible_count,
            "silent_exclusions": 0,
            "by_family": {
                family: {
                    "audited": sum(
                        item["task_family"] == family for item in eligibility
                    ),
                    "eligible": sum(
                        item["task_family"] == family and item["eligible"]
                        for item in eligibility
                    ),
                    "abstain": sum(
                        item["task_family"] == family
                        and item["decision"] == "abstain"
                        for item in eligibility
                    ),
                }
                for family in ("OfficeQA", "SpreadsheetBench")
            },
        },
        "task_eligibility": eligibility,
        "call_and_cost_matrix": matrix,
        "checks": checks,
        "decision": (
            "execution_plan_frozen_paid_execution_closed"
            if all(checks.values())
            else "execution_preflight_failed_paid_execution_closed"
        ),
        "network_calls": 0,
        "paid_api_calls": 0,
        "provider_attempts": 0,
        "scientific_claim_effect": "none_protocol_and_budget_evidence_only",
    }


def write_artifact(
    config: dict[str, Any],
    output_dir: Path,
) -> dict[str, Any]:
    if output_dir.exists():
        raise FileExistsError(
            f"refusing to overwrite immutable artifact: {output_dir}"
        )
    result = audit(deepcopy(config))
    output_dir.mkdir(parents=True)
    result_path = output_dir / "execution_budget_audit.json"
    result_path.write_text(
        json.dumps(result, indent=2, ensure_ascii=True),
        encoding="utf-8",
    )
    manifest = {
        "schema_version": 1,
        "experiment": config["experiment"],
        "config_path": _relative(CONFIG_PATH),
        "config_sha256": sha256_file(CONFIG_PATH),
        "expected_runs": 1,
        "available_runs": 1,
        "complete_grid": all(result["checks"].values()),
        "analysis_only": True,
        "network_calls": 0,
        "paid_api_calls": 0,
        "provider_attempts": 0,
        "runs": [
            {
                "run_id": "phase1q_zero_network_execution_budget_preflight",
                "status": "completed",
                "result_path": _relative(result_path),
                "file_sha256": sha256_file(result_path),
            }
        ],
        "aggregate_fingerprint": sha256_file(result_path),
    }
    (output_dir / "run_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=True),
        encoding="utf-8",
    )
    return result


def main() -> int:
    result = write_artifact(read_json(CONFIG_PATH), OUTPUT_DIR)
    print(json.dumps(result, indent=2, ensure_ascii=True))
    return (
        0
        if result["decision"]
        == "execution_plan_frozen_paid_execution_closed"
        else 1
    )


if __name__ == "__main__":
    raise SystemExit(main())
