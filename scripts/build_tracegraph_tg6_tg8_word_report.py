"""Build the Chinese TG6-TG8 research report as a formatted Word document."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper/acl2027/reports/TRACEGRAPH_TG6_TG8_RESEARCH_REPORT_20260917_ZH.md"
OUTPUT = ROOT / "paper/acl2027/reports/TRACEGRAPH_TG6_TG8_RESEARCH_REPORT_20260917_ZH.docx"

BLUE = "244A73"
TEAL = "2B6F6D"
LIGHT_BLUE = "EAF1F7"
LIGHT_GRAY = "F2F4F6"
MID_GRAY = "D4DAE0"
TEXT = "202A33"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, latin="Aptos", east_asia="宋体", size=None, bold=None, color=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def add_hyperlink(paragraph, text, target):
    relationship = paragraph.part.relate_to(
        target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\)|\\\(.+?\\\))")


def add_inline(paragraph, text):
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, latin="Consolas", east_asia="等线", size=9.5, color="3A4652")
            run.font.highlight_color = None
        elif token.startswith("["):
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, target)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, latin="Cambria Math", east_asia="等线")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    headings = {
        "Title": (22, "黑体", BLUE, 0, 12),
        "Heading 1": (16, "黑体", BLUE, 12, 6),
        "Heading 2": (13, "黑体", TEAL, 10, 4),
        "Heading 3": (11.5, "黑体", TEXT, 8, 3),
    }
    for name, (size, font, color, before, after) in headings.items():
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(2)

    if "LaTeX Formula" not in styles:
        style = styles.add_style("LaTeX Formula", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Cambria Math"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(TEXT)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_together = True


def configure_page(document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    header = section.header.paragraphs[0]
    header.text = "TraceGraph TG6–TG8 研究与审计报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.runs[0], east_asia="等线", size=8.5, color="6E7781")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    set_run_font(run, east_asia="等线", size=8.5, color="6E7781")
    add_field(footer, "PAGE")
    run = footer.add_run(" 页")
    set_run_font(run, east_asia="等线", size=8.5, color="6E7781")


def add_cover(document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(90)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("从合法动作到任务完成")
    set_run_font(run, latin="Aptos Display", east_asia="黑体", size=25, bold=True, color=BLUE)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("TraceGraph TG6–TG8 研究、实验设计与终态审计报告")
    set_run_font(run, latin="Aptos Display", east_asia="黑体", size=17, bold=True, color=TEAL)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(26)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("ACL 2027 研究项目阶段报告")
    set_run_font(run, east_asia="等线", size=12, color="4D5965")

    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.5)
    table.columns[1].width = Cm(10.5)
    values = [
        ("报告日期", "2026 年 9 月 17 日"),
        ("证据范围", "TG6–TG8 本地冻结实验、终态审计与任务聚类分析"),
        ("公式格式", "正文公式保留为 LaTeX 源码，便于论文复用"),
        ("结论级别", "开发试验信号；非确认性结论或跨域通用证明"),
    ]
    for row, (key, value) in zip(table.rows, values):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_margins(cell, 110, 120, 110, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        key_run = row.cells[0].paragraphs[0].add_run(key)
        set_run_font(key_run, east_asia="黑体", size=10, bold=True, color=BLUE)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, east_asia="宋体", size=10)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(65)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("基于完整审计证据生成；原始实验结果保持只读")
    set_run_font(run, east_asia="等线", size=9.5, color="6E7781")
    document.add_page_break()


def add_toc(document):
    heading = document.add_paragraph("目录", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = document.add_paragraph()
    add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u')
    hint = document.add_paragraph("提示：在 Word 中选中目录并按 F9 可更新页码。")
    hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(hint.runs[0], east_asia="等线", size=8.5, color="6E7781")
    document.add_page_break()


def add_design_overview(document):
    document.add_heading("实验设计速览", level=1)
    paragraph = document.add_paragraph()
    add_inline(paragraph, "本页先给出完整研究链条。后文逐阶段说明研究动机、假设、对照、样本、停止规则、指标、结果与证据边界。")
    table = document.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ("阶段", "核心问题", "实验设计", "主要指标", "停止/判定规则", "结果")
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, BLUE)
        run = cell.paragraphs[0].add_run(label)
        set_run_font(run, east_asia="黑体", size=8.5, bold=True, color="FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rows = [
        ("TG6 初始", "区分环境故障与能力失败", "40 个冻结任务；零重试；首个硬违规停止", "环境重置、轨迹合法性、成功", "遇到环境初始化异常即终止", "第 2 个任务 KeyError；不可作能力结论"),
        ("TG6 修复", "原选择器在可执行环境中是否有效", "40 个任务；每条最多 50 步", "成功率、硬不变量", "40 条完整或首个违规停止", "0/40 成功；结构检查通过"),
        ("TG6 诊断", "防循环和探索多样性是否足够", "24 任务 × 3 条件 = 72 条", "成功、二周期、唯一动作/快照", "预注册二周期相对下降 ≥50%", "三组均 0/24；门槛未通过"),
        ("TG7", "目标锚点与进度记忆是否足够", "24 新任务 × 3 条件 = 72 条", "成功、二周期、目标表面命中", "成功提升或二周期相对下降 ≥50%", "三组均 0/24；representation-limited"),
        ("TG8", "表示与控制的主效应和交互", "30 任务 × 3 重复 × 4 条件 = 360 条", "第 50 步完成率与二周期发生率", "任务聚类 bootstrap；区间方向判定 pilot signal", "0/90、0/90、3/90、42/90"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            set_cell_margins(cell, 60, 65, 60, 65)
            if index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            set_run_font(run, east_asia="宋体", size=7.8, bold=index == 0)
    document.add_page_break()


def add_formula(document, lines):
    paragraph = document.add_paragraph(style="LaTeX Formula")
    paragraph.paragraph_format.left_indent = Cm(1.2)
    paragraph.paragraph_format.right_indent = Cm(1.2)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GRAY)
    properties.append(shading)
    run = paragraph.add_run("\\[\n" + "\n".join(lines).strip() + "\n\\]")
    set_run_font(run, latin="Cambria Math", east_asia="等线", size=10.5)


def add_markdown_table(document, lines):
    values = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(values) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in values[1]):
        values.pop(1)
    table = document.add_table(rows=1, cols=len(values[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, values[0]):
        set_cell_shading(cell, BLUE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(text)
        set_run_font(run, east_asia="黑体", size=8.5, bold=True, color="FFFFFF")
    for row_index, row_values in enumerate(values[1:]):
        cells = table.add_row().cells
        for column, (cell, text) in enumerate(zip(cells, row_values)):
            set_cell_margins(cell, 60, 75, 60, 75)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if column == 0 and row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            elif row_index % 2 == 1:
                set_cell_shading(cell, "F8F9FA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body(document, markdown):
    lines = markdown.splitlines()
    index = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(line.strip() for line in paragraph_buffer).replace("  ", " ")
            paragraph = document.add_paragraph()
            add_inline(paragraph, text)
            paragraph_buffer = []

    page_break_headings = {
        "五、TG6：从环境故障到“合法但无效”",
        "六、TG7：记住目标，还不等于会推进目标",
        "七、TG8：设计审查与正式实现",
        "八、TG8 正式运行与审计结果",
        "十、现在能够支持什么，不能支持什么",
        "附录 A：证据索引",
    }
    skipped_h1 = False
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            flush_paragraph()
            index += 1
            continue
        if line.startswith("# ") and not skipped_h1:
            skipped_h1 = True
            index += 1
            continue
        if line.startswith("#"):
            flush_paragraph()
            match = re.match(r"^(#{1,3})\s+(.+)$", line)
            if match:
                level = len(match.group(1))
                title = match.group(2)
                if title in page_break_headings:
                    document.add_page_break()
                document.add_heading(title, level=level)
            index += 1
            continue
        if line.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue
        if line == r"\[":
            flush_paragraph()
            formula = []
            index += 1
            while index < len(lines) and lines[index].strip() != r"\]":
                formula.append(lines[index])
                index += 1
            index += 1
            add_formula(document, formula)
            continue
        if line.startswith("> "):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.right_indent = Cm(0.8)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            properties = paragraph._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:color"), TEAL)
            left.set(qn("w:space"), "8")
            border.append(left)
            properties.append(border)
            add_inline(paragraph, line[2:])
            index += 1
            continue
        ordered = re.match(r"^(\d+)\.\s+(.+)$", line)
        bullet = re.match(r"^-\s+(.+)$", line)
        if ordered or bullet:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number" if ordered else "List Bullet")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, ordered.group(2) if ordered else bullet.group(1))
            index += 1
            continue
        if line == "---":
            flush_paragraph()
            paragraph = document.add_paragraph()
            properties = paragraph._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), MID_GRAY)
            border.append(bottom)
            properties.append(border)
            index += 1
            continue
        paragraph_buffer.append(line.rstrip("  "))
        index += 1
    flush_paragraph()


def prevent_table_row_splitting(document):
    for table in document.tables:
        for row in table.rows:
            properties = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            properties.append(cant_split)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_together = True


def main():
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing Word report: {OUTPUT}")
    document = Document()
    configure_styles(document)
    configure_page(document)
    add_cover(document)
    add_toc(document)
    add_design_overview(document)
    add_body(document, SOURCE.read_text(encoding="utf-8"))
    prevent_table_row_splitting(document)
    properties = document.core_properties
    properties.title = "从合法动作到任务完成：TraceGraph TG6–TG8 研究、实验设计与终态审计报告"
    properties.subject = "ACL 2027 TraceGraph TG6–TG8"
    properties.author = "SummerSkillOpt Research Project"
    properties.keywords = "TraceGraph, TG6, TG7, TG8, ALFWorld, audit, experiment design"
    properties.comments = "Formulas are retained as LaTeX source blocks."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
