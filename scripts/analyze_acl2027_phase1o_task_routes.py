#!/usr/bin/env python3
"""Run ACL 2027 Phase 1O zero-network task-route adjudication."""
from __future__ import annotations

import argparse
import csv
import inspect
import json
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from skillopt.evaluation.cross_task_protocol import (
    apply_spreadsheet_edits,
    target_cells,
    validate_officeqa_trace,
    validate_spreadsheet_edits,
)
try:
    from scripts.run_acl2027_phase1f_smoke import read_json, sha256_file, utc_now, write_json
except ModuleNotFoundError:
    from run_acl2027_phase1f_smoke import read_json, sha256_file, utc_now, write_json

DEFAULT_CONFIG = ROOT / "configs/acl2027/phase1o_task_route_adjudication_v1.json"
DEFAULT_OUTPUT = ROOT / "artifacts/acl2027_phase1o_task_route_adjudication_v1"


class AdjudicationError(RuntimeError):
    """Raised when the frozen zero-network adjudication cannot proceed."""


def _resolve(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def _relative(path: Path) -> str:
    try:
        path = path.relative_to(ROOT)
    except ValueError:
        pass
    return str(path).replace("\\", "/")


def _decimal(value: Any) -> Decimal:
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise AdjudicationError(f"non-numeric deterministic input: {value!r}") from exc


def _json_number(value: Decimal) -> int | str:
    return int(value) if value == value.to_integral_value() else format(value, "f")


def validate_config(config_path: Path = DEFAULT_CONFIG) -> dict[str, Any]:
    config = read_json(config_path)
    if config.get("phase") != "1O":
        raise AdjudicationError("config phase must be 1O")
    execution = config.get("execution", {})
    if not all((
        execution.get("paid_api_allowed") is False,
        execution.get("network_calls_allowed") is False,
        execution.get("provider_calls") == 0,
        execution.get("formal_scaling_allowed") is False,
        execution.get("qwen3_8_max_allowed") is False,
        execution.get("officeqa_24_allowed") is False,
        execution.get("spreadsheetbench_40_allowed") is False,
    )):
        raise AdjudicationError("network, paid, Max, batch, and scaling gates must remain closed")
    for name, item in config.get("immutable_inputs", {}).items():
        path = _resolve(item["path"])
        if not path.is_file() or sha256_file(path) != item["sha256"]:
            raise AdjudicationError(f"immutable input hash mismatch: {name}")
    office = config["constructors"]["officeqa"]
    sheet = config["constructors"]["spreadsheetbench"]
    if office.get("allowed_inputs") != ["officeqa_phase1n_run"]:
        raise AdjudicationError("OfficeQA constructor boundary changed")
    if sheet.get("allowed_inputs") != ["spreadsheetbench_initial_workbook"]:
        raise AdjudicationError("SpreadsheetBench constructor boundary changed")
    if not office.get("reference_answer_forbidden") or not sheet.get("golden_workbook_forbidden"):
        raise AdjudicationError("constructor leakage guard is missing")
    return config


def constructor_leakage_audit() -> dict[str, Any]:
    forbidden = {"reference", "reference_path", "reference_csv", "golden", "golden_path", "golden_workbook"}
    office = sorted(set(inspect.signature(construct_officeqa_from_frozen_run).parameters) & forbidden)
    sheet = sorted(set(inspect.signature(construct_spreadsheet_formulas).parameters) & forbidden)
    return {
        "officeqa_forbidden_parameter_overlap": office,
        "spreadsheetbench_forbidden_parameter_overlap": sheet,
        "passed": not office and not sheet,
    }


def construct_officeqa_from_frozen_run(
    run_path: Path,
    output_path: Path,
    *,
    task_id: str,
    required_time_basis: str,
    expected_period_count: int,
) -> dict[str, Any]:
    run = read_json(run_path)
    if str(run.get("task_id")) != task_id:
        raise AdjudicationError("OfficeQA frozen run task ID mismatch")
    payload = run.get("parsed_response")
    if not isinstance(payload, dict):
        raise AdjudicationError("OfficeQA frozen run has no parsed response")
    trace = validate_officeqa_trace(
        payload,
        required_time_basis=required_time_basis,
        expected_period_count=expected_period_count,
    )
    calculation = payload.get("calculation", {})
    if calculation.get("operation") != "sum":
        raise AdjudicationError("OfficeQA deterministic executor only accepts sum")
    operands = calculation.get("operands")
    if not isinstance(operands, list) or len(operands) != expected_period_count:
        raise AdjudicationError("OfficeQA operand count does not match the contract")
    periods = [str(item.get("period", "")) for item in operands if isinstance(item, dict)]
    if len(periods) != expected_period_count or len(set(periods)) != expected_period_count:
        raise AdjudicationError("OfficeQA operands must contain unique periods")
    evidence = payload.get("evidence", [])
    evidence_values = {
        str(item.get("period")): _decimal(item.get("value"))
        for item in evidence if isinstance(item, dict) and item.get("period") is not None
    }
    operand_values = {str(item["period"]): _decimal(item.get("value")) for item in operands}
    evidence_operand_match = evidence_values == operand_values
    recomputed = sum(operand_values.values(), Decimal(0))
    model_result = _decimal(calculation.get("result"))
    model_answer = _decimal(payload.get("answer"))
    record = {
        "schema_version": 1,
        "task_family": "OfficeQA",
        "task_id": task_id,
        "constructor_inputs": {"frozen_run_sha256": sha256_file(run_path)},
        "reference_accessed": False,
        "trace_validation": trace,
        "operand_count": len(operands),
        "evidence_operand_match": evidence_operand_match,
        "recomputed_answer": _json_number(recomputed),
        "model_answer": _json_number(model_answer),
        "model_calculation_result": _json_number(model_result),
        "model_arithmetic_match": recomputed == model_result == model_answer,
        "execution_decision": "execute" if trace["eligible_for_gate"] and evidence_operand_match else "abstain",
        "route_candidate": "keep_with_deterministic_executor_and_abstain_on_evidence_mismatch",
    }
    write_json(output_path, record)
    return record


def evaluate_officeqa_construction(
    construction_path: Path,
    reference_csv: Path,
    *,
    task_id: str,
) -> dict[str, Any]:
    if not construction_path.is_file():
        raise AdjudicationError("OfficeQA construction must be persisted before evaluation")
    construction = read_json(construction_path)
    with reference_csv.open(newline="", encoding="utf-8-sig") as handle:
        rows = [row for row in csv.DictReader(handle) if str(row.get("uid")) == task_id]
    if len(rows) != 1:
        raise AdjudicationError("OfficeQA reference row did not resolve uniquely")
    reference = _decimal(rows[0]["answer"])
    recomputed = _decimal(construction["recomputed_answer"])
    matched = recomputed == reference
    return {
        "task_family": "OfficeQA",
        "task_id": task_id,
        "construction_sha256": sha256_file(construction_path),
        "reference_answer": _json_number(reference),
        "recomputed_answer": _json_number(recomputed),
        "exact_answer_match": matched,
        "route_decision": (
            "keep_with_deterministic_executor_and_abstain_on_evidence_mismatch"
            if matched and construction["execution_decision"] == "execute" else "replace_or_narrow"
        ),
    }


def construct_spreadsheet_formulas(
    initial_workbook: Path,
    output_workbook: Path,
    record_path: Path,
    *,
    task_id: str,
    target_range: str,
    source_rows: list[int],
    points_per_column: int,
) -> dict[str, Any]:
    min_col, min_row, max_col, max_row = range_boundaries(target_range)
    target_rows = list(range(min_row, max_row + 1))
    if len(target_rows) != len(source_rows):
        raise AdjudicationError("source rows must map one-to-one to target rows")
    workbook = load_workbook(initial_workbook, data_only=False, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        for col in range(min_col, max_col + 1):
            for row in source_rows:
                _decimal(sheet.cell(row=row, column=col).value)
        sheet_name = sheet.title
    finally:
        workbook.close()
    edits = []
    for target_row, source_row in zip(target_rows, source_rows):
        for col in range(min_col, max_col + 1):
            letter = get_column_letter(col)
            source_range = f"{letter}${source_rows[0]}:{letter}${source_rows[-1]}"
            formula = f"=IF({letter}{source_row}=MAX({source_range}),{points_per_column}/COUNTIF({source_range},MAX({source_range})),0)"
            edits.append({"sheet": sheet_name, "cell": f"{letter}{target_row}", "formula": formula})
    payload = {
        "target_range": target_range,
        "edits": edits,
        "explanation": "Per-column maximum with equal division of six points among ties.",
    }
    schema = validate_spreadsheet_edits(payload, expected_target_range=target_range)
    apply_spreadsheet_edits(initial_workbook, output_workbook, payload)
    record = {
        "schema_version": 1,
        "task_family": "SpreadsheetBench",
        "task_id": task_id,
        "constructor_inputs": {"initial_workbook_sha256": sha256_file(initial_workbook)},
        "golden_accessed": False,
        "scoring_contract": {
            "target_range": target_range,
            "source_rows": source_rows,
            "points_per_column": points_per_column,
            "comparison_axis": "within_each_column",
        },
        "payload": payload,
        "schema_validation": schema,
        "generated_workbook": _relative(output_workbook),
        "generated_workbook_sha256": sha256_file(output_workbook),
        "route_candidate": "redesign_as_constrained_executor_backed_task",
    }
    write_json(record_path, record)
    return record


def evaluate_spreadsheet_construction(
    construction_path: Path,
    generated_workbook: Path,
    golden_workbook: Path,
    *,
    task_id: str,
    target_range: str,
) -> dict[str, Any]:
    if not construction_path.is_file() or not generated_workbook.is_file():
        raise AdjudicationError("SpreadsheetBench construction must be persisted before evaluation")
    predicted = load_workbook(generated_workbook, data_only=False, read_only=True)
    golden = load_workbook(golden_workbook, data_only=False, read_only=True)
    try:
        sheet_name = golden.sheetnames[0]
        mismatches = []
        matches = 0
        for cell in target_cells(target_range):
            actual = predicted[sheet_name][cell].value
            expected = golden[sheet_name][cell].value
            if actual == expected:
                matches += 1
            else:
                mismatches.append({"cell": cell, "predicted": actual, "golden": expected})
    finally:
        predicted.close()
        golden.close()
    total = len(target_cells(target_range))
    exact = matches == total
    return {
        "task_family": "SpreadsheetBench",
        "task_id": task_id,
        "construction_sha256": sha256_file(construction_path),
        "generated_workbook_sha256": sha256_file(generated_workbook),
        "exact_formula_matches": matches,
        "target_cells": total,
        "mismatch_count": len(mismatches),
        "mismatches": mismatches,
        "exact_formula_match": exact,
        "route_decision": "redesign_as_constrained_executor_backed_task" if exact else "replace_or_narrow",
    }


def execute_adjudication(
    *,
    config_path: Path = DEFAULT_CONFIG,
    output_dir: Path = DEFAULT_OUTPUT,
) -> dict[str, Any]:
    config = validate_config(config_path)
    if output_dir.exists():
        raise AdjudicationError(f"refusing to overwrite existing artifact: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=False)
    inputs = {name: _resolve(item["path"]) for name, item in config["immutable_inputs"].items()}
    leakage = constructor_leakage_audit()
    if not leakage["passed"]:
        raise AdjudicationError("constructor signature leakage audit failed")

    office_spec = config["constructors"]["officeqa"]
    office_path = output_dir / "officeqa_construction.json"
    construct_officeqa_from_frozen_run(
        inputs["officeqa_phase1n_run"], office_path,
        task_id=office_spec["task_id"],
        required_time_basis=office_spec["required_time_basis"],
        expected_period_count=office_spec["expected_period_count"],
    )
    office_eval = evaluate_officeqa_construction(
        office_path, inputs["officeqa_reference_csv"], task_id=office_spec["task_id"]
    )

    sheet_spec = config["constructors"]["spreadsheetbench"]
    generated = output_dir / "spreadsheetbench_45635_constrained.xlsx"
    sheet_path = output_dir / "spreadsheetbench_construction.json"
    construct_spreadsheet_formulas(
        inputs["spreadsheetbench_initial_workbook"], generated, sheet_path,
        task_id=sheet_spec["task_id"],
        target_range=sheet_spec["target_range"],
        source_rows=sheet_spec["source_rows"],
        points_per_column=sheet_spec["points_per_column"],
    )
    sheet_eval = evaluate_spreadsheet_construction(
        sheet_path, generated, inputs["spreadsheetbench_golden_workbook"],
        task_id=sheet_spec["task_id"], target_range=sheet_spec["target_range"],
    )

    records = [office_eval, sheet_eval]
    scientific_path = output_dir / "scientific_results.json"
    write_json(scientific_path, records)
    runs = []
    for index, record in enumerate(records, start=1):
        run_path = output_dir / f"run_{index:02d}.json"
        write_json(run_path, record)
        runs.append({
            "run_id": f"phase1o_route_{index:02d}_{record['task_family'].lower()}",
            "task_family": record["task_family"],
            "result_path": _relative(run_path),
            "file_sha256": sha256_file(run_path),
            "route_decision": record["route_decision"],
        })
    passed = office_eval["exact_answer_match"] and sheet_eval["exact_formula_match"]
    manifest = {
        "schema_version": 1,
        "experiment": config["experiment"],
        "artifact_role": "zero_network_construct_first_route_adjudication",
        "status": "completed",
        "created_at": utc_now(),
        "completed_at": utc_now(),
        "config_path": _relative(config_path),
        "config_sha256": sha256_file(config_path),
        "immutable_input_hashes": {name: item["sha256"] for name, item in config["immutable_inputs"].items()},
        "constructor_leakage_audit": leakage,
        "expected_runs": 2,
        "available_runs": 2,
        "complete_grid": True,
        "runs": runs,
        "network_calls": 0,
        "paid_api_calls": 0,
        "provider_attempts": 0,
        "formal_scaling": False,
        "decision": "task_routes_retained_with_executor_backing" if passed else "task_route_replacement_or_narrowing_required",
        "scientific_scope": config["decision_gate"]["scientific_scope"],
        "scientific_results_path": _relative(scientific_path),
        "aggregate_fingerprint": sha256_file(scientific_path),
        "artifact_hashes": {
            "officeqa_construction": sha256_file(office_path),
            "spreadsheetbench_construction": sha256_file(sheet_path),
            "spreadsheetbench_generated_workbook": sha256_file(generated),
        },
    }
    write_json(output_dir / "run_manifest.json", manifest)
    return manifest


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    manifest = execute_adjudication(config_path=args.config, output_dir=args.output_dir)
    print(json.dumps(manifest, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
