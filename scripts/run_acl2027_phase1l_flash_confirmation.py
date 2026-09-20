#!/usr/bin/env python3
"""Run the authorized ACL 2027 Phase 1L two-call Flash confirmation."""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openai import OpenAI
from openpyxl import load_workbook

from skillopt.evaluation.cross_task_protocol import (
    apply_spreadsheet_edits,
    parse_json_object_preserving_raw,
    target_cells,
    validate_officeqa_trace,
    validate_spreadsheet_edits,
)
try:
    from scripts.run_acl2027_phase1f_smoke import (
        _raw_response,
        _usage,
        canonical_hash,
        read_json,
        sha256_file,
        utc_now,
        write_json,
    )
except ModuleNotFoundError:
    from run_acl2027_phase1f_smoke import (
        _raw_response,
        _usage,
        canonical_hash,
        read_json,
        sha256_file,
        utc_now,
        write_json,
    )

DEFAULT_CONFIG = ROOT / "configs/acl2027/phase1l_flash_live_confirmation_v1.json"
DEFAULT_STATE = ROOT / "paper/acl2027/experiment_state.json"
DEFAULT_OUTPUT = ROOT / "artifacts/acl2027_phase1l_flash_live_confirmation_v1"


class SmokeError(RuntimeError):
    """Raised when the bounded Phase 1L confirmation cannot proceed safely."""


def _relative(path: Path) -> str:
    try:
        value = path.relative_to(ROOT)
    except ValueError:
        value = path
    return str(value).replace("\\", "/")


def validate_inputs(config_path: Path = DEFAULT_CONFIG) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    config = read_json(config_path)
    if config.get("phase") != "1L":
        raise SmokeError("config phase must be 1L")
    auth = config.get("authorization", {})
    execution = config.get("execution", {})
    if auth.get("authorized") is not True or auth.get("authorized_logical_calls") != 2:
        raise SmokeError("Phase 1L authorization must cover exactly two logical calls")
    if auth.get("max_provider_attempts") != 2 or auth.get("explicit_retries") != 0:
        raise SmokeError("Phase 1L provider attempt limit changed")
    closed = ("qwen3.7_plus_authorized", "qwen3.8_max_authorized", "officeqa_24_authorized", "spreadsheetbench_40_authorized", "formal_scaling_authorized")
    if any(auth.get(key) is not False for key in closed):
        raise SmokeError("Phase 1L model escalation, batch, and scaling authorization must remain closed")
    if execution.get("paid_api_allowed") is not True or execution.get("network_calls_allowed") is not True:
        raise SmokeError("Phase 1L live confirmation permission is missing")
    if execution.get("formal_scaling_allowed") is not False or execution.get("sdk_max_retries") != 0:
        raise SmokeError("Phase 1L scaling or retry policy changed")

    frozen = config["frozen_preflight"]
    for path_key, hash_key in (("config_path", "config_sha256"), ("request_plan_path", "request_plan_sha256")):
        if sha256_file(ROOT / frozen[path_key]) != frozen[hash_key]:
            raise SmokeError(f"frozen preflight {path_key} hash mismatch")
    full_plan = read_json(ROOT / frozen["request_plan_path"])
    indices = frozen.get("selected_request_indices", [])
    if indices != [1, 2] or len(full_plan) != 4:
        raise SmokeError("Phase 1K request selection changed")
    plan = [full_plan[index - 1] for index in indices]
    if [row.get("task_family") for row in plan] != frozen["task_order"]:
        raise SmokeError("frozen task order mismatch")
    if [str(row.get("task_id")) for row in plan] != frozen["task_ids"]:
        raise SmokeError("frozen task IDs mismatch")
    for row in plan:
        family = row["task_family"]
        if row.get("model") != frozen["model_id"] or row.get("request", {}).get("model") != frozen["model_id"]:
            raise SmokeError(f"{family} model is not frozen qwen3.6-flash")
        if row.get("request_hash") != frozen["request_hashes"][family]:
            raise SmokeError(f"{family} frozen request hash mismatch")
        if canonical_hash(row["request"]) != row["request_hash"]:
            raise SmokeError(f"{family} request content changed after Phase 1K")
    return config, plan


def assert_live_allowed(config: dict[str, Any], state_path: Path = DEFAULT_STATE) -> None:
    state = read_json(state_path)
    policy = state.get("execution_policy", {})
    phase = state.get("current_phase", {})
    if policy.get("paid_api_allowed") is not True:
        raise PermissionError("repository paid permission is closed")
    if policy.get("formal_scaling_allowed") is not False:
        raise PermissionError("formal scaling must remain disabled")
    if phase.get("id") != "1L" or phase.get("status") != "in_progress":
        raise PermissionError("repository state is not Phase 1L in_progress")
    reason = str(policy.get("reason", "")).lower()
    if "exactly two" not in reason or "qwen3.6-flash" not in reason or "no retries" not in reason or "24" not in reason or "40" not in reason:
        raise PermissionError("repository authorization is not narrowly scoped to the two-call Flash confirmation")
    if config["execution"].get("batch_closed") is not True:
        raise PermissionError("batch gate is open")


def _number(value: Any) -> Decimal | None:
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError):
        return None


def _office_reference_answer(config: dict[str, Any]) -> str:
    path = ROOT / config["validation"]["officeqa_reference_answer_source"]
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = [row for row in csv.DictReader(handle) if str(row.get("uid")) == "UID0001"]
    if len(rows) != 1:
        raise SmokeError("OfficeQA reference row did not resolve uniquely")
    return str(rows[0]["answer"])


def _spreadsheet_paths(config: dict[str, Any]) -> tuple[Path, Path]:
    task_dir = ROOT / config["validation"]["spreadsheetbench_task_dir"]
    initial = sorted(task_dir.glob("*_init.xlsx"))
    golden = sorted(task_dir.glob("*_golden.xlsx"))
    if len(initial) != 1 or len(golden) != 1:
        raise SmokeError("SpreadsheetBench initial/golden workbook did not resolve uniquely")
    return initial[0], golden[0]


def _formula_comparison(predicted: Path, golden: Path, target_range: str) -> dict[str, Any]:
    pred_wb = load_workbook(predicted, data_only=False, read_only=True)
    gold_wb = load_workbook(golden, data_only=False, read_only=True)
    try:
        sheet = gold_wb.sheetnames[0]
        mismatches = []
        for cell in target_cells(target_range):
            actual = pred_wb[sheet][cell].value
            expected = gold_wb[sheet][cell].value
            if actual != expected:
                mismatches.append({"cell": cell, "predicted": actual, "golden": expected})
        return {"exact_formula_match": not mismatches, "mismatch_count": len(mismatches), "mismatches": mismatches}
    finally:
        pred_wb.close()
        gold_wb.close()


def _validate_task(config: dict[str, Any], family: str, parsed: dict[str, Any], output_dir: Path) -> dict[str, Any]:
    validation = config["validation"]
    if family == "OfficeQA":
        trace = validate_officeqa_trace(
            parsed,
            required_time_basis=validation["officeqa_required_time_basis"],
            expected_period_count=validation["officeqa_expected_period_count"],
        )
        reference = _office_reference_answer(config)
        answer_match = _number(parsed.get("answer")) == _number(reference)
        result_match = _number(parsed.get("calculation", {}).get("result")) == _number(reference)
        return {
            **trace,
            "answer_matches_reference": answer_match,
            "calculation_result_matches_reference": result_match,
            "task_valid": trace["eligible_for_gate"] and answer_match and result_match,
        }

    schema = validate_spreadsheet_edits(parsed, expected_target_range=validation["spreadsheetbench_expected_target_range"])
    initial, golden = _spreadsheet_paths(config)
    output_workbook = output_dir / "spreadsheetbench_45635_predicted.xlsx"
    apply_spreadsheet_edits(initial, output_workbook, parsed)
    formulas = _formula_comparison(output_workbook, golden, validation["spreadsheetbench_expected_target_range"])
    return {
        **schema,
        **formulas,
        "workbook_persisted": output_workbook.is_file(),
        "output_workbook": _relative(output_workbook),
        "task_valid": schema["exact_coverage"] and schema["formulas_valid"] and formulas["exact_formula_match"],
    }


def execute_smoke(
    *,
    config_path: Path = DEFAULT_CONFIG,
    state_path: Path = DEFAULT_STATE,
    output_dir: Path = DEFAULT_OUTPUT,
    live: bool = False,
    client: Any | None = None,
) -> dict[str, Any]:
    config, plan = validate_inputs(config_path)
    if not live:
        return {
            "mode": "dry-run",
            "logical_calls": 0,
            "planned_calls": 2,
            "config_sha256": sha256_file(config_path),
            "requests": [{"task_family": row["task_family"], "task_id": row["task_id"], "request_hash": row["request_hash"]} for row in plan],
        }
    assert_live_allowed(config, state_path)
    api_key = os.environ.get(config["execution"]["api_key_env"], "")
    if client is None and not api_key.strip():
        raise PermissionError(f"missing {config['execution']['api_key_env']}")
    if output_dir.exists():
        raise SmokeError(f"refusing to overwrite existing artifact: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=False)

    manifest = {
        "schema_version": 1,
        "experiment": config["experiment"],
        "artifact_role": "authorized_cost_aware_two_call_flash_confirmation_only",
        "status": "in_progress",
        "started_at": utc_now(),
        "config_path": _relative(config_path),
        "config_sha256": sha256_file(config_path),
        "runner_sha256": sha256_file(Path(__file__)),
        "provider": {"endpoint": config["execution"]["endpoint"], "model_id": config["execution"]["model_id"], "sdk_max_retries": 0},
        "limits": {"planned_logical_calls": 2, "max_provider_attempts": 2, "explicit_retries": 0, "batch_authorized": False},
        "request_plan_sha256": canonical_hash(plan),
        "calls": [],
    }
    write_json(output_dir / "request_plan.json", plan)
    write_json(output_dir / "run_manifest.json", manifest)
    provider = client or OpenAI(
        api_key=api_key,
        base_url=config["execution"]["endpoint"].rstrip("/"),
        timeout=float(config["execution"]["timeout_seconds"]),
        max_retries=0,
    )

    records: list[dict[str, Any]] = []
    results_path = output_dir / "results.jsonl"
    for index, item in enumerate(plan, start=1):
        started = time.perf_counter()
        record: dict[str, Any] = {
            "call_index": index,
            "task_family": item["task_family"],
            "task_id": item["task_id"],
            "request_hash": item["request_hash"],
            "attempt_count": 1,
            "started_at": utc_now(),
        }
        try:
            request = item["request"]
            response = provider.chat.completions.create(
                model=request["model"],
                messages=request["messages"],
                max_tokens=request["max_tokens"],
                temperature=request["temperature"],
                extra_body={"enable_thinking": request["enable_thinking"]},
            )
            raw_response = _raw_response(response)
            choices = getattr(response, "choices", None) or []
            raw_text = str(getattr(choices[0].message, "content", "")) if choices else ""
            input_tokens, output_tokens, total_tokens = _usage(response)
            record.update({
                "usage_known": True,
                "input_tokens": input_tokens,
                "output_tokens": output_tokens,
                "total_tokens": total_tokens,
                "provider_request_id": str(getattr(response, "id", "") or ""),
                "raw_response": raw_response,
                "raw_text": raw_text,
            })
            normalized = parse_json_object_preserving_raw(raw_text)
            task_validation = _validate_task(config, item["task_family"], normalized.value, output_dir)
            record.update({
                "status": "success",
                "parsed_response": normalized.value,
                "normalizations": list(normalized.normalizations),
                "task_validation": task_validation,
                "task_valid": task_validation["task_valid"],
            })
        except Exception as exc:
            record.update({
                "status": "failed",
                "task_valid": False,
                "error_type": type(exc).__name__,
                "error": str(exc),
            })
            record.setdefault("usage_known", False)
        record["duration_seconds"] = round(time.perf_counter() - started, 6)
        records.append(record)
        with results_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(record, ensure_ascii=True, sort_keys=True) + "\n")
        manifest["calls"].append({
            "call_index": index,
            "task_family": item["task_family"],
            "task_id": item["task_id"],
            "status": record["status"],
            "task_valid": record["task_valid"],
            "usage_known": record["usage_known"],
        })
        write_json(output_dir / "run_manifest.json", manifest)

    scientific_records = []
    for record in records:
        scientific_records.append({key: value for key, value in record.items() if key not in {"started_at", "duration_seconds", "provider_request_id", "raw_response"}})
    scientific_path = output_dir / "scientific_results.json"
    write_json(scientific_path, scientific_records)
    run_entries = []
    for record in scientific_records:
        result_path = output_dir / f"run_{int(record['call_index']):02d}.json"
        write_json(result_path, record)
        run_entries.append({
            "run_id": f"phase1l_flash_{int(record['call_index']):02d}_{str(record['task_family']).lower()}",
            "result_path": _relative(result_path),
            "file_sha256": sha256_file(result_path),
            "task_family": record["task_family"],
            "status": record["status"],
            "task_valid": record["task_valid"],
        })
    usage_known = all(record.get("usage_known") is True for record in records)
    usage = {
        "input_tokens": sum(int(record.get("input_tokens", 0)) for record in records),
        "output_tokens": sum(int(record.get("output_tokens", 0)) for record in records),
        "total_tokens": sum(int(record.get("total_tokens", 0)) for record in records),
        "exact_for_successful_calls": usage_known,
    }
    passed = len(records) == 2 and usage_known and all(record.get("status") == "success" and record.get("task_valid") is True for record in records)
    manifest.update({
        "status": "completed",
        "runs": run_entries,
        "completed_at": utc_now(),
        "expected_runs": 2,
        "available_runs": len(records),
        "complete_grid": len(records) == 2,
        "provider_attempts": len(records),
        "usage": usage,
        "decision": "flash_task_valid_floor_passed_plus_not_needed" if passed else "flash_task_valid_floor_failed_plus_review_required",
        "scientific_results_path": _relative(scientific_path),
        "aggregate_fingerprint": sha256_file(scientific_path),
        "network_calls": len(records),
        "paid_api_calls": len(records),
    })
    write_json(output_dir / "run_manifest.json", manifest)
    return manifest


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--live-smoke", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    result = execute_smoke(live=args.live_smoke)
    print(json.dumps(result, indent=2, ensure_ascii=True))
    return 0 if not args.live_smoke or result["decision"] == "flash_task_valid_floor_passed_plus_not_needed" else 1


if __name__ == "__main__":
    raise SystemExit(main())

