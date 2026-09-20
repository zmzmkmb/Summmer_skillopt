#!/usr/bin/env python3
"""Run the ACL 2027 Phase 1F two-call cross-task smoke.

Dry-run is the default. Live execution requires --live-smoke, repository
permission, the frozen Phase 1F config, and a fresh output directory.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openai import OpenAI
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "configs/acl2027/phase1f_real_cross_task_triage_pilot_v1.json"
DEFAULT_STATE = ROOT / "paper/acl2027/experiment_state.json"
DEFAULT_OUTPUT = ROOT / "artifacts/acl2027_phase1f_real_cross_task_triage_pilot_v1/bounded_smoke_20260810"


class SmokeError(RuntimeError):
    """Raised when the bounded smoke cannot continue safely."""


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonical_hash(value: Any) -> str:
    payload = json.dumps(value, sort_keys=True, ensure_ascii=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, indent=2, ensure_ascii=True, sort_keys=True) + "\n", encoding="utf-8")


def validate_inputs(config_path: Path) -> dict[str, Any]:
    config = read_json(config_path)
    execution = config.get("execution", {})
    if config.get("phase") != "1F":
        raise SmokeError("config phase must be 1F")
    if execution.get("default_mode") != "dry-run":
        raise SmokeError("Phase 1F must remain dry-run by default")
    if execution.get("formal_scaling_allowed") is not False:
        raise SmokeError("formal scaling must remain disabled")
    if execution.get("batch_requires_separate_review") is not True:
        raise SmokeError("batch review gate is missing")
    if int(config.get("pairing", {}).get("smoke_logical_calls", -1)) != 2:
        raise SmokeError("bounded smoke must contain exactly two logical calls")
    for family, task in config["data"].items():
        manifest = ROOT / task["source_manifest"]
        if sha256_file(manifest) != task["source_manifest_sha256"]:
            raise SmokeError(f"{family} split manifest hash mismatch")
        if family == "OfficeQA":
            payload = ROOT / task["payload_csv"]
            expected = task["payload_csv_sha256"]
        else:
            payload = ROOT / task["payload_dataset"]
            expected = task["payload_dataset_sha256"]
        if sha256_file(payload) != expected:
            raise SmokeError(f"{family} payload hash mismatch")
    return config


def assert_live_allowed(config: dict[str, Any], state_path: Path) -> None:
    state = read_json(state_path)
    policy = state.get("execution_policy", {})
    phase = state.get("current_phase", {})
    if config.get("execution", {}).get("paid_api_allowed") is not True:
        raise PermissionError("live smoke is disabled by the immutable config")
    if policy.get("paid_api_allowed") is not True:
        raise PermissionError("live smoke is disabled by repository state")
    if policy.get("formal_scaling_allowed") is not False:
        raise PermissionError("formal scaling must remain disabled")
    if phase.get("id") != "1F" or phase.get("status") != "in_progress":
        raise PermissionError("repository state is not Phase 1F in_progress")
    reason = str(policy.get("reason", ""))
    if "two-call bounded smoke" not in reason or "does not authorize the 24/40 batch" not in reason:
        raise PermissionError("repository permission is not narrowly scoped to the two-call smoke")


def _office_row(config: dict[str, Any]) -> dict[str, str]:
    task = config["data"]["OfficeQA"]
    target = str(task["smoke_ids"][0])
    with (ROOT / task["payload_csv"]).open(newline="", encoding="utf-8-sig") as handle:
        rows = [row for row in csv.DictReader(handle) if str(row.get("uid", row.get("UID", ""))) == target]
    if len(rows) != 1:
        raise SmokeError(f"OfficeQA smoke UID {target} did not resolve uniquely")
    return rows[0]


def _office_context(source_name: str) -> tuple[Path, str]:
    transformed = ROOT / "data/officeqa_verified/treasury_bulletins_parsed/transformed"
    matches = list(transformed.glob(source_name))
    if len(matches) != 1:
        raise SmokeError(f"OfficeQA source {source_name} did not resolve uniquely")
    source = matches[0]
    lines = source.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    heading = "Budget Expenditures Classified as General, by Major Functions"
    try:
        start = next(i for i, line in enumerate(lines) if heading in line)
        end = next(i for i in range(start + 1, len(lines)) if lines[i].startswith("Source:"))
    except StopIteration as exc:
        raise SmokeError("OfficeQA evidence table could not be extracted deterministically") from exc
    context = "\n".join(lines[start : end + 1]).strip()
    if "National defense" not in context or "1940-January" not in context:
        raise SmokeError("OfficeQA evidence table failed content checks")
    return source, context


def _spreadsheet_item(config: dict[str, Any]) -> dict[str, Any]:
    task = config["data"]["SpreadsheetBench"]
    target = str(task["smoke_ids"][0])
    payload = read_json(ROOT / task["payload_dataset"])
    rows = [row for row in payload if str(row.get("id")) == target]
    if len(rows) != 1:
        raise SmokeError(f"SpreadsheetBench smoke ID {target} did not resolve uniquely")
    return rows[0]


def _workbook_snapshot(task_dir: Path) -> tuple[Path, dict[str, Any]]:
    init_files = sorted(task_dir.glob("*_init.xlsx"))
    if not init_files:
        candidate = task_dir / "initial.xlsx"
        init_files = [candidate] if candidate.exists() else []
    if len(init_files) != 1:
        raise SmokeError(f"expected one initial workbook in {task_dir}")
    workbook_path = init_files[0]
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    cells: dict[str, Any] = {}
    for row in sheet.iter_rows(min_row=1, max_row=8, min_col=1, max_col=6):
        for cell in row:
            if cell.value is not None:
                cells[cell.coordinate] = cell.value
    snapshot = {"sheet": sheet.title, "cells_A1_F8": cells}
    workbook.close()
    return workbook_path, snapshot


def build_smoke_plan(config: dict[str, Any]) -> list[dict[str, Any]]:
    office = _office_row(config)
    office_source, office_context = _office_context(str(office["source_files"]))
    office_messages = [
        {"role": "system", "content": "Answer from the supplied source only. Return one JSON object with keys answer and evidence. No markdown."},
        {"role": "user", "content": f"Question: {office['question']}\n\nSource excerpt:\n{office_context}"},
    ]
    spreadsheet = _spreadsheet_item(config)
    task_dir = ROOT / "data/spreadsheetbench_verified_400" / str(spreadsheet["spreadsheet_path"])
    workbook_path, snapshot = _workbook_snapshot(task_dir)
    spreadsheet_messages = [
        {"role": "system", "content": "Solve the spreadsheet formula task. Return one JSON object with keys formula, target_range, and explanation. No markdown. Do not claim the workbook was executed."},
        {"role": "user", "content": f"Instruction: {spreadsheet['instruction']}\n\nWorkbook snapshot: {json.dumps(snapshot, ensure_ascii=True, sort_keys=True)}"},
    ]
    common = {"model": config["execution"]["model_id"], "temperature": 0, "enable_thinking": False}
    return [
        {
            "task_family": "OfficeQA", "task_id": str(office["uid"]),
            "messages": office_messages, "max_tokens": 256,
            "source_path": str(office_source.relative_to(ROOT)).replace("\\", "/"),
            "source_sha256": sha256_file(office_source),
            "request_hash": canonical_hash({**common, "messages": office_messages, "max_tokens": 256}),
        },
        {
            "task_family": "SpreadsheetBench", "task_id": str(spreadsheet["id"]),
            "messages": spreadsheet_messages, "max_tokens": 512,
            "source_path": str(workbook_path.relative_to(ROOT)).replace("\\", "/"),
            "source_sha256": sha256_file(workbook_path),
            "request_hash": canonical_hash({**common, "messages": spreadsheet_messages, "max_tokens": 512}),
        },
    ]


def _parse_json_object(text: str) -> dict[str, Any]:
    cleaned = text.strip()
    fence = chr(96) * 3
    if cleaned.startswith(fence):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\s*```$", "", cleaned)
    parsed = json.loads(cleaned)
    if not isinstance(parsed, dict):
        raise SmokeError("provider output must be a JSON object")
    return parsed


def _validate_parsed(family: str, parsed: dict[str, Any]) -> None:
    required = {"answer", "evidence"} if family == "OfficeQA" else {"formula", "target_range", "explanation"}
    missing = sorted(key for key in required if not str(parsed.get(key, "")).strip())
    if missing:
        raise SmokeError(f"{family} output omitted required fields: {missing}")


def _raw_response(response: Any) -> dict[str, Any]:
    dumper = getattr(response, "model_dump", None)
    if callable(dumper):
        return dumper(mode="json")
    choice = response.choices[0]
    return {
        "id": getattr(response, "id", ""),
        "model": getattr(response, "model", ""),
        "choices": [{"message": {"content": getattr(choice.message, "content", "")}}],
        "usage": {
            "prompt_tokens": getattr(response.usage, "prompt_tokens", None),
            "completion_tokens": getattr(response.usage, "completion_tokens", None),
            "total_tokens": getattr(response.usage, "total_tokens", None),
        },
    }


def _usage(response: Any) -> tuple[int, int, int]:
    usage = getattr(response, "usage", None)
    values = [getattr(usage, name, None) for name in ("prompt_tokens", "completion_tokens", "total_tokens")]
    if any(value is None for value in values):
        raise SmokeError("provider response omitted complete usage metadata")
    input_tokens, output_tokens, total_tokens = (int(value) for value in values)
    if min(input_tokens, output_tokens, total_tokens) < 0:
        raise SmokeError("provider usage contains negative values")
    if total_tokens != input_tokens + output_tokens:
        raise SmokeError("provider usage identity failed")
    return input_tokens, output_tokens, total_tokens


def execute_smoke(*, config_path: Path = DEFAULT_CONFIG, state_path: Path = DEFAULT_STATE,
                  output_dir: Path = DEFAULT_OUTPUT, live: bool = False,
                  client: Any | None = None) -> dict[str, Any]:
    config = validate_inputs(config_path)
    plan = build_smoke_plan(config)
    if len(plan) != 2 or [item["task_family"] for item in plan] != ["OfficeQA", "SpreadsheetBench"]:
        raise SmokeError("smoke plan is not the frozen two-family order")
    if not live:
        return {
            "mode": "dry-run", "logical_calls": 0, "planned_calls": 2,
            "config_sha256": sha256_file(config_path),
            "requests": [{key: item[key] for key in ("task_family", "task_id", "request_hash", "max_tokens", "source_sha256")} for item in plan],
        }
    assert_live_allowed(config, state_path)
    api_key = os.environ.get(str(config["execution"]["api_key_env"]), "")
    if client is None and not api_key.strip():
        raise PermissionError(f"missing {config['execution']['api_key_env']}")
    if output_dir.exists():
        raise SmokeError(f"refusing to overwrite existing artifact: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=False)
    manifest = {
        "schema_version": 1, "experiment": config["experiment"],
        "artifact_role": "two_call_bounded_smoke_only", "status": "in_progress",
        "started_at": utc_now(),
        "config_path": str(config_path.relative_to(ROOT)).replace("\\", "/"),
        "config_sha256": sha256_file(config_path), "runner_sha256": sha256_file(Path(__file__)),
        "provider": {"endpoint": config["execution"]["endpoint"], "model_id": config["execution"]["model_id"], "sdk_max_retries": 0},
        "limits": {"planned_logical_calls": 2, "max_provider_attempts": 2, "explicit_retries": 0, "batch_authorized": False},
        "request_plan_sha256": canonical_hash(plan), "calls": [],
    }
    write_json(output_dir / "request_plan.json", plan)
    write_json(output_dir / "run_manifest.json", manifest)
    provider = client or OpenAI(api_key=api_key, base_url=str(config["execution"]["endpoint"]).rstrip("/"), timeout=90.0, max_retries=0)
    results_path = output_dir / "results.jsonl"
    error: Exception | None = None
    for index, item in enumerate(plan, start=1):
        started = time.perf_counter()
        record: dict[str, Any] = {
            "call_index": index, "task_family": item["task_family"], "task_id": item["task_id"],
            "request_hash": item["request_hash"], "attempt_count": 1, "started_at": utc_now(),
        }
        known_usage: tuple[int, int, int] | None = None
        raw_response_payload: dict[str, Any] | None = None
        raw_text = ""
        try:
            response = provider.chat.completions.create(
                model=config["execution"]["model_id"], messages=item["messages"],
                max_tokens=item["max_tokens"], temperature=0, extra_body={"enable_thinking": False},
            )
            input_tokens, output_tokens, total_tokens = _usage(response)
            known_usage = (input_tokens, output_tokens, total_tokens)
            raw_response_payload = _raw_response(response)
            choices = getattr(response, "choices", None) or []
            if not choices or not str(getattr(choices[0].message, "content", "")).strip():
                raise SmokeError("provider returned empty content")
            raw_text = str(choices[0].message.content)
            parsed = _parse_json_object(raw_text)
            _validate_parsed(item["task_family"], parsed)
            record.update({
                "status": "success", "usage_known": True,
                "input_tokens": input_tokens, "output_tokens": output_tokens, "total_tokens": total_tokens,
                "provider_request_id": str(getattr(response, "id", "") or ""),
                "raw_response": raw_response_payload, "raw_text": raw_text, "parsed_response": parsed,
            })
        except Exception as exc:
            record.update({"status": "failed", "error_type": type(exc).__name__, "error": str(exc)})
            if known_usage is None:
                record["usage_known"] = False
            else:
                record.update({
                    "usage_known": True,
                    "input_tokens": known_usage[0],
                    "output_tokens": known_usage[1],
                    "total_tokens": known_usage[2],
                    "raw_response": raw_response_payload,
                    "raw_text": raw_text,
                })
            error = exc
        record["latency_seconds"] = round(time.perf_counter() - started, 6)
        record["finished_at"] = utc_now()
        with results_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(record, ensure_ascii=True, sort_keys=True) + "\n")
        manifest["calls"].append({key: record.get(key) for key in (
            "call_index", "task_family", "task_id", "request_hash", "status", "usage_known",
            "input_tokens", "output_tokens", "total_tokens", "attempt_count", "latency_seconds", "error_type", "error")})
        write_json(output_dir / "run_manifest.json", manifest)
        if error is not None:
            break
    successful = [call for call in manifest["calls"] if call["status"] == "success"]
    manifest["finished_at"] = utc_now()
    manifest["provider_attempts"] = len(manifest["calls"])
    manifest["successful_calls"] = len(successful)
    manifest["usage"] = {
        "input_tokens": sum(int(call.get("input_tokens") or 0) for call in successful),
        "output_tokens": sum(int(call.get("output_tokens") or 0) for call in successful),
        "total_tokens": sum(int(call.get("total_tokens") or 0) for call in successful),
        "exact_for_successful_calls": all(call.get("usage_known") is True for call in successful),
    }
    manifest["status"] = "passed" if len(successful) == 2 and error is None else "stopped"
    manifest["decision"] = "smoke_passed_batch_still_closed" if manifest["status"] == "passed" else "smoke_stopped_batch_closed"
    if results_path.exists():
        manifest["results_sha256"] = sha256_file(results_path)
    write_json(output_dir / "run_manifest.json", manifest)
    if error is not None:
        raise SmokeError(f"bounded smoke stopped after call {len(manifest['calls'])}: {type(error).__name__}: {error}") from error
    return manifest


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--state", type=Path, default=DEFAULT_STATE)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--live-smoke", action="store_true", help="Issue exactly two live calls; dry-run is the default.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        result = execute_smoke(config_path=args.config.resolve(), state_path=args.state.resolve(), output_dir=args.out_dir.resolve(), live=args.live_smoke)
    except (SmokeError, PermissionError) as exc:
        print(f"ERROR: {exc}")
        return 1
    print(json.dumps(result, indent=2, ensure_ascii=True, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
